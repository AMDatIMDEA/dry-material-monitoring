from __future__ import annotations

from copy import deepcopy
from dataclasses import replace
from datetime import datetime, timedelta, timezone
import json
from pathlib import Path
import sys

from openpyxl import load_workbook
import pytest
import yaml


PACKAGE_ROOT = Path(__file__).resolve().parents[1]
REPOSITORY_ROOT = PACKAGE_ROOT.parent
sys.path.insert(0, str(REPOSITORY_ROOT))
sys.path.insert(0, str(PACKAGE_ROOT))

from experiment_records import COMMON_COLUMNS, make_measurement_id, sha256_file
from material_volume.config import load_config
from material_volume.pipeline import MaterialVolumePipeline
from material_volume.session import (
    DepthMeasurementRequest,
    PreparedDepthMethod,
    StorageProfile,
)
from material_volume.synthetic import SyntheticDepthSource


TRIGGER = datetime(2026, 8, 13, 10, 45, 12, 345678, tzinfo=timezone.utc)


@pytest.fixture(scope="module")
def fast_config_and_calibration():
    config = deepcopy(load_config(PACKAGE_ROOT / "config.yaml"))
    config.camera.depth_width = 160
    config.camera.depth_height = 90
    config.camera.enable_color = True
    # Freeze scientific geometry inside the fixture so production config edits
    # cannot change the 40 mm == 50% record assertions below.
    config.tube.usable_height_mm = 80.0
    config.calibration.center_x_px = None
    config.calibration.center_y_px = None
    config.detection_roi.center_x_px = None
    config.detection_roi.center_y_px = None
    config.detection_roi.frame_width_px = None
    config.detection_roi.frame_height_px = None
    config.fusion.burst_frames = 7
    # The reduced synthetic image keeps this suite fast; use a proportionate
    # coverage gate so it exercises accepted common records, not metrology.
    config.reconstruction.minimum_surface_coverage = 0.15
    with SyntheticDepthSource(config, fill_percent=0.0) as empty_source:
        empty_burst = empty_source.capture_burst(config.fusion.burst_frames)
    calibration = MaterialVolumePipeline.calibrate(empty_burst, config)
    return config, calibration


def request(
    output_root: Path | None,
    experiment_id: str,
    profile: str,
    *,
    index: int | None = None,
    measurement_id: str | None = None,
    trigger: datetime = TRIGGER,
) -> DepthMeasurementRequest:
    return DepthMeasurementRequest(
        experiment_id=experiment_id,
        purpose="hardware-free session test",
        operator_notes="résultat synthétique",
        material_name="test polymer",
        total_capacity_ml=None,
        bulk_density_g_per_ml=0.5,
        total_possible_weight_g=None,
        reference_material_weight_g=10.0,
        manual_material_level_mm=40.0,
        output_root=output_root,
        storage_profile=profile,
        measurement_index=index,
        measurement_id=measurement_id,
        trigger_time_utc=trigger,
        capture_start_utc=trigger + timedelta(milliseconds=10),
        capture_end_utc=trigger + timedelta(milliseconds=900),
        processing_time_utc=trigger + timedelta(seconds=1),
        acquisition_mode="synthetic",
        software_repository=REPOSITORY_ROOT,
    )


def run_session(config, calibration, session_request, fill_percent=35.0):
    with SyntheticDepthSource(config, fill_percent=fill_percent) as source:
        return PreparedDepthMethod(config, calibration, source).run_one(session_request)


class TestDepthResearchSession:
    def test_config_without_research_block_keeps_backward_compatible_defaults(
        self,
        tmp_path: Path,
    ) -> None:
        data = yaml.safe_load((PACKAGE_ROOT / "config.yaml").read_text(encoding="utf-8"))
        data.pop("research")
        legacy_config = tmp_path / "legacy config.yaml"
        legacy_config.write_text(
            yaml.safe_dump(data, sort_keys=False),
            encoding="utf-8",
        )
        loaded = load_config(legacy_config)
        assert loaded.research.storage_profile == "research"
        assert loaded.research_output_path == tmp_path / "research_records"

    @pytest.mark.parametrize(
        ("profile", "present", "absent"),
        [
            (
                "compact",
                {"result.json", "height_map.png", "depth_diagnostics.json", "artifact_manifest.json"},
                {"height_map_mm.npy", "material_surface_mm.ply", "color_snapshot.png"},
            ),
            (
                "research",
                {
                    "result.json",
                    "height_map.png",
                    "height_map_mm.npy",
                    "material_surface_mm.ply",
                    "color_snapshot.png",
                    "depth_diagnostics.json",
                    "artifact_manifest.json",
                },
                {"raw_depth/raw_depth_burst_z16.npz"},
            ),
            (
                "full_raw",
                {
                    "result.json",
                    "height_map.png",
                    "height_map_mm.npy",
                    "material_surface_mm.ply",
                    "color_snapshot.png",
                    "depth_diagnostics.json",
                    "artifact_manifest.json",
                    "raw_depth/raw_depth_burst_z16.npz",
                },
                set(),
            ),
        ],
    )
    def test_storage_profiles_and_provenance(
        self,
        tmp_path: Path,
        fast_config_and_calibration,
        profile: str,
        present: set[str],
        absent: set[str],
    ) -> None:
        config, calibration = fast_config_and_calibration
        output_root = tmp_path / "Étude avec espaces"
        outcome = run_session(
            config,
            calibration,
            request(output_root, f"profile-{profile}", profile),
        )
        assert outcome.estimate is not None
        files = {
            path.relative_to(outcome.artifact_directory).as_posix()
            for path in outcome.artifact_directory.rglob("*")
            if path.is_file()
        }
        assert present <= files
        assert not (absent & files)

        session = json.loads(
            (outcome.session_directory / "session_manifest.json").read_text(encoding="utf-8")
        )
        assert session["storage_profile"] == profile
        assert session["scientifically_validated"] is False
        config_snapshot = outcome.session_directory / session["effective_config"]
        assert sha256_file(config_snapshot) == session["config_sha256"]
        calibration_reference = outcome.session_directory / session["calibration_reference"]
        assert calibration_reference.is_file()
        assert sha256_file(calibration_reference) == session["calibration_sha256"]
        assert (outcome.session_directory / "provenance" / "environment.json").is_file()

        artifact_manifest = json.loads(outcome.artifact_manifest.read_text(encoding="utf-8"))
        for artifact in artifact_manifest["artifacts"]:
            path = outcome.artifact_directory / artifact["path"]
            assert sha256_file(path) == artifact["sha256"]

    def test_default_unicode_root_auto_indexes_and_common_workbook(
        self,
        tmp_path: Path,
        fast_config_and_calibration,
    ) -> None:
        base_config, calibration = fast_config_and_calibration
        config = deepcopy(base_config)
        config.research.directory = str((tmp_path / "Données de recherche").resolve())
        first = run_session(config, calibration, request(None, "session-auto", "research"))
        second_trigger = TRIGGER + timedelta(seconds=2)
        second = run_session(
            config,
            calibration,
            request(None, "session-auto", "research", trigger=second_trigger),
            fill_percent=60.0,
        )
        assert first.record.measurement_index == 1
        assert second.record.measurement_index == 2
        assert first.record.measurement_id.startswith("000001_taking_")
        assert second.record.measurement_id.startswith("000002_taking_")
        assert first.session_directory.parent == config.research_output_path

        workbook = first.session_directory / "depth_measurements.xlsx"
        csv_mirror = first.session_directory / "depth_measurements.csv"
        sheet = load_workbook(workbook)["measurements"]
        assert tuple(cell.value for cell in sheet[1]) == COMMON_COLUMNS
        assert sheet.max_row == 3
        assert csv_mirror.is_file()
        records = {
            sheet.cell(row=row, column=COMMON_COLUMNS.index("measurement_id") + 1).value
            for row in (2, 3)
        }
        assert records == {first.record.measurement_id, second.record.measurement_id}
        assert first.record.reference_volume_from_weight_ml == 20.0
        assert first.record.manual_material_percent == 50.0
        assert first.record.reference_volume_from_manual_ml == pytest.approx(
            first.record.total_capacity_ml * 0.5
        )

        diagnostics = json.loads(
            (first.artifact_directory / "depth_diagnostics.json").read_text(encoding="utf-8")
        )
        assert diagnostics["quality"]["heuristic_variability_is_gum_uncertainty"] is False
        assert "surface_coverage" in diagnostics["quality"]
        assert diagnostics["capture"]["camera_serial_number"] == "synthetic"
        assert diagnostics["capture"]["active_stream_profile"]["format"] == "Z16"
        assert diagnostics["capture"]["depth_scale_m"] == 0.001
        assert "yolo" not in json.dumps(diagnostics).lower()

    def test_supplied_shared_identity_and_timestamps(
        self,
        tmp_path: Path,
        fast_config_and_calibration,
    ) -> None:
        config, calibration = fast_config_and_calibration
        shared_id = make_measurement_id(12, TRIGGER)
        outcome = run_session(
            config,
            calibration,
            request(
                tmp_path.resolve(),
                "external-id",
                "compact",
                index=12,
                measurement_id=shared_id,
            ),
        )
        assert outcome.record.measurement_id == shared_id
        assert outcome.record.measurement_index == 12
        assert outcome.record.trigger_time_utc == "2026-08-13T10:45:12.345678Z"
        assert outcome.record.capture_start_utc == "2026-08-13T10:45:12.355678Z"
        assert outcome.record.capture_end_utc == "2026-08-13T10:45:13.245678Z"
        assert outcome.record.processing_time_utc == "2026-08-13T10:45:13.345678Z"

    def test_non_raw_profile_accepts_legacy_source_capture_signature(
        self,
        tmp_path: Path,
        fast_config_and_calibration,
    ) -> None:
        config, calibration = fast_config_and_calibration

        class LegacySource:
            def __init__(self) -> None:
                self.delegate = SyntheticDepthSource(config, fill_percent=35.0)

            def capture_burst(self, frame_count: int):
                return self.delegate.capture_burst(frame_count)

        outcome = PreparedDepthMethod(config, calibration, LegacySource()).run_one(
            request(tmp_path.resolve(), "legacy-source-api", "research")
        )
        assert outcome.estimate is not None

    def test_supplied_index_cannot_be_reused_with_a_new_trigger(
        self,
        tmp_path: Path,
        fast_config_and_calibration,
    ) -> None:
        config, calibration = fast_config_and_calibration
        root = tmp_path.resolve()
        run_session(
            config,
            calibration,
            request(root, "unique-index", "compact", index=7),
        )
        with pytest.raises(ValueError, match="already assigned"):
            run_session(
                config,
                calibration,
                request(
                    root,
                    "unique-index",
                    "compact",
                    index=7,
                    trigger=TRIGGER + timedelta(seconds=1),
                ),
            )

    def test_session_wrapper_does_not_change_synthetic_volume(
        self,
        tmp_path: Path,
        fast_config_and_calibration,
    ) -> None:
        config, calibration = fast_config_and_calibration
        with SyntheticDepthSource(config, fill_percent=35.0) as source:
            baseline = MaterialVolumePipeline(config, calibration).measure(
                source.capture_burst(config.fusion.burst_frames)
            )
        outcome = run_session(
            config,
            calibration,
            request(tmp_path.resolve(), "numeric-regression", "compact"),
        )
        assert outcome.estimate is not None
        assert outcome.estimate.result.fill_percent == baseline.result.fill_percent
        assert outcome.estimate.result.material_volume_ml == baseline.result.material_volume_ml
        assert outcome.estimate.result.empty_volume_ml == baseline.result.empty_volume_ml
        assert outcome.record.estimated_material_volume_ml == baseline.result.material_volume_ml

    def test_relative_custom_output_root_is_rejected(
        self,
        fast_config_and_calibration,
    ) -> None:
        config, calibration = fast_config_and_calibration
        with pytest.raises(ValueError, match="absolute"):
            run_session(
                config,
                calibration,
                request(Path("relative output"), "bad-root", "compact"),
            )

    def test_preallocated_synchronized_reservation_keeps_external_capture_manifest(
        self,
        tmp_path: Path,
        fast_config_and_calibration,
    ) -> None:
        config, calibration = fast_config_and_calibration
        synchronized = request(
            tmp_path.resolve(),
            "prepared-synchronized",
            "compact",
        )
        synchronized = replace(
            synchronized,
            acquisition_mode="synchronized",
            capture_start_utc=None,
            capture_end_utc=None,
            processing_time_utc=None,
        )
        with SyntheticDepthSource(config, fill_percent=35.0) as source:
            prepared = PreparedDepthMethod(config, calibration, source)
            reservation = prepared.reserve(synchronized)
            capture_manifest = reservation.measurement_directory / "capture_manifest.json"
            capture_manifest.write_text(
                json.dumps({"state": "ORCHESTRATOR_OWNS_THIS", "c920": {"frames": []}}),
                encoding="utf-8",
            )
            captured = prepared.capture(
                replace(
                    synchronized,
                    measurement_index=reservation.measurement_index,
                    measurement_id=reservation.measurement_id,
                ),
                now_utc=lambda: TRIGGER,
            )
            outcome = prepared.process(
                captured,
                replace(
                    synchronized,
                    measurement_index=reservation.measurement_index,
                    measurement_id=reservation.measurement_id,
                ),
                reservation=reservation,
                manage_capture_manifest=False,
                now_utc=lambda: TRIGGER + timedelta(seconds=2),
            )
        assert outcome.record.measurement_id == reservation.measurement_id
        assert json.loads(capture_manifest.read_text(encoding="utf-8"))["state"] == (
            "ORCHESTRATOR_OWNS_THIS"
        )
        assert (reservation.session_directory / "depth_measurements.xlsx").is_file()
