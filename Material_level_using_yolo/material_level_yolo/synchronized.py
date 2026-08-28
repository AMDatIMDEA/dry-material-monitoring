"""Manifest-led offline YOLO processing for Run_Experiment captures."""

from __future__ import annotations

from contextlib import contextmanager
from dataclasses import dataclass
from datetime import datetime, timezone
import json
import os
from pathlib import Path
import shutil
import tempfile
import time
from typing import Any, Callable, Iterator, Mapping, Protocol

from openpyxl import load_workbook
import yaml

from experiment_records import (
    COMMON_COLUMNS,
    CommonMeasurementRecord,
    MeasurementWorkbookStore,
    RecordStatus,
    format_utc,
    sha256_file,
)

from .config import ProjectConfig
from .domain import ModelProfile
from .errors import YoloFoundationError
from .inference import UltralyticsInferenceAdapter
from .workflow import (
    InferenceAdapterProtocol,
    OfflineMeasurementRequest,
    effective_profile_config_sha256,
    process_saved_images,
)


COMPLETED_STATUSES = frozenset(
    {RecordStatus.COMPLETE_VALID.value, RecordStatus.COMPLETE_INVALID.value}
)


class SynchronizedProcessingError(YoloFoundationError):
    """A synchronized capture cannot be safely processed or reconciled."""


class AdapterFactoryProtocol(Protocol):
    def __call__(self, profile: ModelProfile) -> InferenceAdapterProtocol: ...


@dataclass(frozen=True, slots=True)
class SynchronizedCaptureGroup:
    session_directory: Path
    measurement_directory: Path
    capture_manifest_path: Path
    experiment_id: str
    measurement_id: str
    measurement_index: int
    trigger_time_utc: str
    source_paths: tuple[Path, ...]
    source_hashes: tuple[str, ...]
    capture_start_utc: str
    capture_end_utc: str
    planned_count: int
    capture_state: str


@dataclass(frozen=True, slots=True)
class SynchronizedGroupResult:
    experiment_id: str
    measurement_id: str
    measurement_index: int
    action: str
    profile_name: str
    status: str
    valid: bool | None
    accepted_image_count: int | None
    rejected_image_count: int | None
    processing_revision: int
    artifact_directory: Path


@dataclass(frozen=True, slots=True)
class SynchronizedProcessingReport:
    target: Path
    session_directory: Path
    groups: tuple[SynchronizedGroupResult, ...]


def discover_synchronized_groups(
    target: str | Path,
) -> tuple[SynchronizedCaptureGroup, ...]:
    """Discover one selected measurement or every manifest in a session."""
    selected = Path(target).expanduser().resolve()
    if not selected.is_dir():
        raise SynchronizedProcessingError(
            f"Session or measurement directory not found: {selected}"
        )
    if (selected / "session_manifest.json").is_file():
        session = selected
        manifest_paths = tuple(
            sorted(
                (session / "measurements").glob("*/capture_manifest.json"),
                key=lambda item: item.parent.name,
            )
        )
    elif (selected / "capture_manifest.json").is_file():
        if selected.parent.name != "measurements":
            raise SynchronizedProcessingError(
                "A selected measurement directory must be directly beneath measurements/."
            )
        session = selected.parent.parent.resolve()
        if not (session / "session_manifest.json").is_file():
            raise SynchronizedProcessingError(
                f"Session manifest not found above selected measurement: {selected}"
            )
        manifest_paths = (selected / "capture_manifest.json",)
    else:
        raise SynchronizedProcessingError(
            f"Directory is neither a synchronized session nor measurement: {selected}"
        )
    if not manifest_paths:
        raise SynchronizedProcessingError(f"No capture manifests found beneath: {session}")
    session_manifest = _read_json(session / "session_manifest.json")
    experiment_id = _required_text(session_manifest, "experiment_id", "session manifest")
    groups = tuple(
        _capture_group(session, path.resolve(), experiment_id) for path in manifest_paths
    )
    return tuple(sorted(groups, key=lambda item: item.measurement_index))


def process_synchronized_captures(
    project: ProjectConfig,
    target: str | Path,
    *,
    profile_name: str | None = None,
    force_reprocess: bool = False,
    start_measurement_index: int | None = None,
    adapter_factory: AdapterFactoryProtocol | None = None,
    now_utc: Callable[[], datetime] | None = None,
    software_repository: Path | None = None,
    progress_callback: Callable[[int, int, SynchronizedCaptureGroup], None] | None = None,
) -> SynchronizedProcessingReport:
    """Process selected synchronized captures without acquiring either camera."""
    groups = _groups_from_index(
        discover_synchronized_groups(target),
        start_measurement_index,
    )
    clock = now_utc or (lambda: datetime.now(timezone.utc))
    create_adapter = adapter_factory or (lambda profile: UltralyticsInferenceAdapter(profile))
    adapters: dict[str, InferenceAdapterProtocol] = {}

    def cached_adapter(profile: ModelProfile) -> InferenceAdapterProtocol:
        # A session normally uses one material profile. Keep its CPU model loaded
        # across measurement groups instead of paying the load cost for every six images.
        if profile.name not in adapters:
            adapters[profile.name] = create_adapter(profile)
        return adapters[profile.name]

    results = []
    session = groups[0].session_directory
    for position, group in enumerate(groups, start=1):
        if group.session_directory != session:
            raise SynchronizedProcessingError("Discovered capture groups span sessions.")
        if progress_callback is not None:
            progress_callback(position, len(groups), group)
        results.append(
            _process_group(
                project,
                group,
                profile_name=profile_name,
                force_reprocess=force_reprocess,
                adapter_factory=cached_adapter,
                clock=clock,
                software_repository=software_repository,
            )
        )
    return SynchronizedProcessingReport(
        target=Path(target).expanduser().resolve(),
        session_directory=session,
        groups=tuple(results),
    )


def _groups_from_index(
    groups: tuple[SynchronizedCaptureGroup, ...],
    start_measurement_index: int | None,
) -> tuple[SynchronizedCaptureGroup, ...]:
    if start_measurement_index is None:
        return groups
    if (
        isinstance(start_measurement_index, bool)
        or not isinstance(start_measurement_index, int)
        or start_measurement_index < 1
    ):
        raise SynchronizedProcessingError(
            "start_measurement_index must be an integer of at least 1."
        )
    selected = tuple(
        group for group in groups if group.measurement_index >= start_measurement_index
    )
    if not selected:
        raise SynchronizedProcessingError(
            f"No capture groups have measurement_index >= {start_measurement_index}."
        )
    return selected


def _process_group(
    project: ProjectConfig,
    group: SynchronizedCaptureGroup,
    *,
    profile_name: str | None,
    force_reprocess: bool,
    adapter_factory: AdapterFactoryProtocol,
    clock: Callable[[], datetime],
    software_repository: Path | None,
) -> SynchronizedGroupResult:
    session = group.session_directory
    measurement = group.measurement_directory
    store = MeasurementWorkbookStore(session / "yolo_measurements.xlsx", method="yolo")
    depth_before = _depth_hashes(session)
    journal_path = measurement / "yolo_processing_manifest.json"
    with _exclusive_lock(measurement / ".offline_yolo.lock"):
        try:
            _verify_sources(group)
            row = _load_original_row(store, group)
            profile = project.select_profile_for_material(row.material_name, profile_name)
            reconciled = _reconcile_committed_attempt(
                group,
                row,
                profile,
                journal_path,
            )
            if reconciled is not None and not force_reprocess:
                committed = _load_original_row(store, group)
                _verify_workbook_contract(session, store)
                _finalize_bookkeeping(
                    group,
                    committed,
                    profile,
                    reconciled.processing_revision,
                    None,
                    clock,
                )
                _write_journal(
                    journal_path,
                    group,
                    state="COMPLETED",
                    profile=profile.name,
                    revision=reconciled.processing_revision,
                    status=committed.status,
                    valid=committed.valid,
                    artifact_directory=committed.artifact_directory,
                    recovered_after_interruption=True,
                    updated_utc=format_utc(clock()),
                )
                return reconciled
            if _rollback_incomplete_revision(group, journal_path, clock):
                row = _load_original_row(store, group)
            if row.status in COMPLETED_STATUSES and not force_reprocess:
                _verify_completed_evidence(session, row)
                _verify_workbook_contract(session, store)
                return _result_from_record(
                    row,
                    profile.name,
                    "verified_existing",
                    _recorded_revision(journal_path),
                    session,
                )

            if row.status not in COMPLETED_STATUSES and force_reprocess:
                raise SynchronizedProcessingError(
                    "--force-reprocess is only needed for a completed YOLO row; "
                    "pending and partial rows resume automatically."
                )
            profile.require_frozen_roi()
            _recover_or_archive_residual(group, journal_path, row)
            revision = _next_revision(measurement)
            adapter = adapter_factory(profile)
            _write_journal(
                journal_path,
                group,
                state="VALIDATING",
                profile=profile.name,
                revision=revision,
                updated_utc=format_utc(clock()),
            )
            adapter.load()
            model_hash = adapter.weights_sha256
            if model_hash is None:
                raise SynchronizedProcessingError(
                    "Validated inference adapter did not provide a model SHA-256."
                )
            config_hash = effective_profile_config_sha256(project, profile)
            revision_directory = None
            if row.status in COMPLETED_STATUSES:
                intended_revision = (
                    measurement / "yolo_revisions" / f"revision_{revision:04d}"
                )
                _write_journal(
                    journal_path,
                    group,
                    state="ARCHIVING",
                    profile=profile.name,
                    revision=revision,
                    model_sha256=model_hash,
                    config_sha256=config_hash,
                    recovery_directory=intended_revision.relative_to(session).as_posix(),
                    rollback_required=True,
                    updated_utc=format_utc(clock()),
                )
                revision_directory = _archive_previous_result(
                    group,
                    row,
                    revision,
                    profile,
                    model_hash,
                    config_hash,
                    clock,
                )
            active = measurement / "yolo"
            if active.exists() and any(active.iterdir()):
                _archive_interrupted_evidence(measurement, active, revision, clock)
            _write_journal(
                journal_path,
                group,
                state="PROCESSING",
                profile=profile.name,
                revision=revision,
                model_sha256=model_hash,
                config_sha256=config_hash,
                recovery_directory=(
                    None
                    if revision_directory is None
                    else revision_directory.relative_to(session).as_posix()
                ),
                rollback_required=revision_directory is not None,
                updated_utc=format_utc(clock()),
            )
            session_manifest = _read_json(session / "session_manifest.json")
            usable_height = _usable_internal_height(session, row)
            request = OfflineMeasurementRequest(
                experiment_id=group.experiment_id,
                purpose=session_manifest.get("purpose"),
                total_capacity_ml=float(row.total_capacity_ml),
                output_root=session.parent,
                measurement_index=group.measurement_index,
                measurement_id=group.measurement_id,
                trigger_time_utc=group.trigger_time_utc,
                material_name=row.material_name,
                bulk_density_g_per_ml=row.bulk_density_g_per_ml,
                total_possible_weight_g=row.total_possible_weight_g,
                reference_material_weight_g=row.reference_material_weight_g,
                manual_material_level_mm=row.manual_material_level_mm,
                usable_internal_height_mm=usable_height,
                notes=row.notes,
                acquisition_mode="synchronized",
                capture_start_utc=group.capture_start_utc,
                capture_end_utc=group.capture_end_utc,
                resume_pending_synchronized=True,
                allow_provenance_revision=row.status in COMPLETED_STATUSES,
                processing_revision=revision if row.status in COMPLETED_STATUSES else None,
                source_artifact=group.capture_manifest_path.relative_to(session).as_posix(),
                software_repository=software_repository,
            )
            try:
                outcome = process_saved_images(
                    project,
                    profile.name,
                    group.source_paths,
                    request,
                    adapter=adapter,
                    now_utc=clock,
                )
            except Exception:
                if revision_directory is not None:
                    _rollback_revision(group, revision_directory, clock)
                _write_journal(
                    journal_path,
                    group,
                    state="ERROR",
                    profile=profile.name,
                    revision=revision,
                    rollback_required=False,
                    updated_utc=format_utc(clock()),
                )
                raise
            _verify_completed_evidence(session, outcome.record)
            _verify_workbook_contract(session, store)
            _finalize_bookkeeping(
                group,
                outcome.record,
                profile,
                revision,
                revision_directory,
                clock,
            )
            _write_journal(
                journal_path,
                group,
                state="COMPLETED",
                profile=profile.name,
                revision=revision,
                model_sha256=outcome.record.calibration_or_model_sha256,
                config_sha256=outcome.record.config_sha256,
                status=outcome.record.status,
                valid=outcome.record.valid,
                accepted_image_count=outcome.group.accepted_count,
                rejected_image_count=outcome.group.rejected_count,
                artifact_directory=outcome.record.artifact_directory,
                updated_utc=format_utc(clock()),
            )
            return SynchronizedGroupResult(
                experiment_id=group.experiment_id,
                measurement_id=group.measurement_id,
                measurement_index=group.measurement_index,
                action=("reprocessed_revision" if revision_directory else "processed"),
                profile_name=profile.name,
                status=str(outcome.record.status),
                valid=outcome.record.valid,
                accepted_image_count=outcome.group.accepted_count,
                rejected_image_count=outcome.group.rejected_count,
                processing_revision=revision,
                artifact_directory=outcome.artifact_directory,
            )
        finally:
            _assert_depth_unchanged(session, depth_before)


def _capture_group(
    session: Path,
    manifest_path: Path,
    session_experiment_id: str,
) -> SynchronizedCaptureGroup:
    try:
        manifest_path.relative_to(session)
    except ValueError as exc:
        raise SynchronizedProcessingError(
            f"Capture manifest escapes session: {manifest_path}"
        ) from exc
    data = _read_json(manifest_path)
    experiment_id = _required_text(data, "experiment_id", str(manifest_path))
    measurement_id = _required_text(data, "measurement_id", str(manifest_path))
    if experiment_id != session_experiment_id:
        raise SynchronizedProcessingError(
            f"Capture/session experiment IDs differ: {manifest_path}"
        )
    measurement_index = data.get("measurement_index")
    if isinstance(measurement_index, bool) or not isinstance(measurement_index, int) or measurement_index < 1:
        raise SynchronizedProcessingError(f"Invalid measurement_index in {manifest_path}")
    if manifest_path.parent.name != measurement_id:
        raise SynchronizedProcessingError(
            f"Measurement directory and manifest ID differ: {manifest_path}"
        )
    c920 = data.get("c920")
    if not isinstance(c920, Mapping):
        raise SynchronizedProcessingError(f"Missing c920 mapping in {manifest_path}")
    frames = c920.get("frames")
    if not isinstance(frames, list) or not frames:
        raise SynchronizedProcessingError(
            f"Capture has no persisted C920 frames to process: {manifest_path}"
        )
    planned = c920.get("planned_count")
    if isinstance(planned, bool) or not isinstance(planned, int) or planned < 1:
        raise SynchronizedProcessingError(f"Invalid C920 planned_count in {manifest_path}")
    paths: list[Path] = []
    hashes: list[str] = []
    timestamps: list[str] = []
    indexes: list[int] = []
    for frame in frames:
        if not isinstance(frame, Mapping):
            raise SynchronizedProcessingError(f"Invalid C920 frame entry in {manifest_path}")
        index = frame.get("index")
        if isinstance(index, bool) or not isinstance(index, int) or index < 1:
            raise SynchronizedProcessingError(f"Invalid C920 frame index in {manifest_path}")
        relative = Path(_required_text(frame, "path", "C920 frame"))
        if relative.is_absolute():
            raise SynchronizedProcessingError("C920 frame paths must be session-relative.")
        resolved = (session / relative).resolve()
        try:
            resolved.relative_to(session.resolve())
        except ValueError as exc:
            raise SynchronizedProcessingError(
                f"C920 frame path escapes session: {relative}"
            ) from exc
        indexes.append(index)
        paths.append(resolved)
        hashes.append(_required_text(frame, "sha256", "C920 frame"))
        timestamps.append(_required_text(frame, "actual_utc", "C920 frame"))
    if len(set(indexes)) != len(indexes) or indexes != sorted(indexes):
        raise SynchronizedProcessingError(f"C920 frame indexes are duplicate/unsorted: {manifest_path}")
    if len(paths) > planned:
        raise SynchronizedProcessingError(f"C920 frame count exceeds planned_count: {manifest_path}")
    state = _required_text(data, "state", str(manifest_path))
    if state == "SAVED" and len(paths) != planned:
        raise SynchronizedProcessingError(
            f"SAVED capture does not contain its planned C920 frame count: {manifest_path}"
        )
    return SynchronizedCaptureGroup(
        session_directory=session,
        measurement_directory=manifest_path.parent,
        capture_manifest_path=manifest_path,
        experiment_id=experiment_id,
        measurement_id=measurement_id,
        measurement_index=measurement_index,
        trigger_time_utc=_required_text(data, "trigger_time_utc", str(manifest_path)),
        source_paths=tuple(paths),
        source_hashes=tuple(hashes),
        capture_start_utc=timestamps[0],
        capture_end_utc=timestamps[-1],
        planned_count=planned,
        capture_state=state,
    )


def _verify_sources(group: SynchronizedCaptureGroup) -> None:
    for path, expected in zip(group.source_paths, group.source_hashes):
        if not path.is_file():
            raise SynchronizedProcessingError(f"Captured C920 source image is missing: {path}")
        actual = sha256_file(path)
        if actual != expected:
            raise SynchronizedProcessingError(
                f"Captured C920 source hash mismatch for {path}: expected {expected}, found {actual}"
            )


def _load_original_row(
    store: MeasurementWorkbookStore,
    group: SynchronizedCaptureGroup,
) -> CommonMeasurementRecord:
    matches = [
        row
        for row in store.load_records()
        if row.experiment_id == group.experiment_id
        and row.measurement_id == group.measurement_id
        and row.method == "yolo"
    ]
    if len(matches) != 1:
        raise SynchronizedProcessingError(
            f"Expected exactly one pending/completed YOLO row for {group.measurement_id}; "
            f"found {len(matches)}."
        )
    row = matches[0]
    checks = {
        "measurement_index": (row.measurement_index, group.measurement_index),
        "trigger_time_utc": (format_utc(row.trigger_time_utc), format_utc(group.trigger_time_utc)),
        "capture_start_utc": (format_utc(row.capture_start_utc), format_utc(group.capture_start_utc)),
        "capture_end_utc": (format_utc(row.capture_end_utc), format_utc(group.capture_end_utc)),
    }
    mismatches = [name for name, values in checks.items() if values[0] != values[1]]
    if mismatches:
        raise SynchronizedProcessingError(
            "YOLO row and capture manifest differ for: " + ", ".join(mismatches)
        )
    return row


def _verify_completed_evidence(session: Path, record: CommonMeasurementRecord) -> None:
    if record.status not in COMPLETED_STATUSES:
        raise SynchronizedProcessingError("Only completed YOLO rows have verifiable result evidence.")
    artifact = _session_path(session, record.artifact_directory, "artifact_directory")
    manifest_path = artifact / "artifact_manifest.json"
    manifest = _read_json(manifest_path)
    if manifest.get("measurement_id") != record.measurement_id or manifest.get("method") != "yolo":
        raise SynchronizedProcessingError(f"YOLO artifact manifest identity mismatch: {manifest_path}")
    entries = manifest.get("artifacts")
    if not isinstance(entries, list) or not entries:
        raise SynchronizedProcessingError(f"YOLO artifact manifest has no entries: {manifest_path}")
    for entry in entries:
        if not isinstance(entry, Mapping):
            raise SynchronizedProcessingError(f"Invalid artifact entry: {manifest_path}")
        path = _artifact_child(artifact, _required_text(entry, "path", "artifact entry"))
        if not path.is_file():
            raise SynchronizedProcessingError(f"YOLO artifact is missing: {path}")
        if path.stat().st_size != entry.get("bytes") or sha256_file(path) != entry.get("sha256"):
            raise SynchronizedProcessingError(f"YOLO artifact integrity check failed: {path}")


def _verify_workbook_contract(session: Path, store: MeasurementWorkbookStore) -> None:
    if not store.verify_mirror():
        raise SynchronizedProcessingError("YOLO Excel/CSV mirror verification failed.")
    depth = session / "depth_measurements.xlsx"
    yolo = session / "yolo_measurements.xlsx"
    if depth.is_file():
        books = []
        try:
            books = [load_workbook(depth, read_only=True), load_workbook(yolo, read_only=True)]
            headers = [tuple(cell.value for cell in book["measurements"][1]) for book in books]
            if headers[0] != headers[1] or headers[0] != COMMON_COLUMNS:
                raise SynchronizedProcessingError(
                    "Depth and YOLO primary-sheet columns are not identical."
                )
        finally:
            for book in books:
                book.close()


def _usable_internal_height(session: Path, row: CommonMeasurementRecord) -> float | None:
    if row.manual_material_level_mm is None:
        return None
    for path in (
        session / "effective_config.yaml",
        session / "provenance" / "depth_effective_config.yaml",
    ):
        if not path.is_file():
            continue
        data = yaml.safe_load(path.read_text(encoding="utf-8-sig")) or {}
        tube = data.get("tube") if isinstance(data, Mapping) else None
        value = tube.get("usable_height_mm") if isinstance(tube, Mapping) else None
        if isinstance(value, (int, float)) and not isinstance(value, bool) and float(value) > 0:
            return float(value)
    raise SynchronizedProcessingError(
        "Manual level metadata exists, but usable tube height is unavailable from the "
        "frozen depth configuration."
    )


def _reconcile_committed_attempt(
    group: SynchronizedCaptureGroup,
    row: CommonMeasurementRecord,
    profile: ModelProfile,
    journal_path: Path,
) -> SynchronizedGroupResult | None:
    if not journal_path.is_file():
        return None
    journal = _read_json(journal_path)
    if journal.get("state") != "PROCESSING":
        return None
    current = _load_original_row(
        MeasurementWorkbookStore(group.session_directory / "yolo_measurements.xlsx", method="yolo"),
        group,
    )
    if current.status not in COMPLETED_STATUSES:
        return None
    if (
        current.config_sha256 != journal.get("config_sha256")
        or current.calibration_or_model_sha256 != journal.get("model_sha256")
    ):
        return None
    try:
        _verify_completed_evidence(group.session_directory, current)
    except SynchronizedProcessingError:
        return None
    return _result_from_record(
        current,
        profile.name,
        "recovered_committed",
        int(journal.get("processing_revision") or 1),
        group.session_directory,
    )


def _rollback_incomplete_revision(
    group: SynchronizedCaptureGroup,
    journal_path: Path,
    clock: Callable[[], datetime],
) -> bool:
    if not journal_path.is_file():
        return False
    journal = _read_json(journal_path)
    if journal.get("state") not in {"ARCHIVING", "PROCESSING"} or not journal.get(
        "rollback_required"
    ):
        return False
    relative = journal.get("recovery_directory")
    if not isinstance(relative, str) or not relative:
        return False
    directory = _session_path(group.session_directory, relative, "recovery_directory")
    if not directory.is_dir() or not (directory / "revision.json").is_file():
        # The process stopped before any prior evidence moved; the existing row
        # and active directory remain authoritative.
        return False
    _rollback_revision(group, directory, clock)
    _write_journal(
        journal_path,
        group,
        state="ERROR",
        profile=journal.get("profile"),
        revision=journal.get("processing_revision"),
        rollback_required=False,
        recovered_after_interruption=True,
        updated_utc=format_utc(clock()),
    )
    return True


def _recover_or_archive_residual(
    group: SynchronizedCaptureGroup,
    journal_path: Path,
    row: CommonMeasurementRecord,
) -> None:
    active = group.measurement_directory / "yolo"
    if not active.exists():
        return
    if row.status in COMPLETED_STATUSES:
        return
    revision = _next_revision(group.measurement_directory)
    _archive_interrupted_evidence(
        group.measurement_directory,
        active,
        revision,
        lambda: datetime.now(timezone.utc),
    )
    if journal_path.is_file():
        journal = _read_json(journal_path)
        journal["residual_evidence_archived"] = True
        _atomic_json(journal_path, journal)


def _archive_previous_result(
    group: SynchronizedCaptureGroup,
    row: CommonMeasurementRecord,
    revision: int,
    profile: ModelProfile,
    model_hash: str,
    config_hash: str,
    clock: Callable[[], datetime],
) -> Path:
    session = group.session_directory
    measurement = group.measurement_directory
    directory = measurement / "yolo_revisions" / f"revision_{revision:04d}"
    directory.mkdir(parents=True, exist_ok=False)
    active = measurement / "yolo"
    if not active.is_dir():
        raise SynchronizedProcessingError(
            f"Completed YOLO row has no active evidence directory: {active}"
        )
    _atomic_json(directory / "previous_record.json", row.to_mapping())
    _atomic_json(
        directory / "revision.json",
        {
            "revision_schema_version": 1,
            "state": "PROCESSING_REPLACEMENT",
            "processing_revision": revision,
            "archived_utc": format_utc(clock()),
            "previous_config_sha256": row.config_sha256,
            "previous_model_sha256": row.calibration_or_model_sha256,
            "requested_profile": profile.name,
            "requested_config_sha256": config_hash,
            "requested_model_sha256": model_hash,
        },
    )
    recovery = directory / "recovery"
    for source in (
        session / "yolo_measurements.xlsx",
        session / "yolo_measurements.csv",
        session / "session_manifest.json",
        session / "provenance" / f"yolo_effective_config_{profile.name}.yaml",
    ):
        if source.is_file():
            _copy_file(source, recovery / source.name)
    # The move is last: before this point the old workbook and active evidence
    # remain mutually consistent. Once moved, all rollback inputs are durable.
    _archive_directory(active, directory / "evidence")
    return directory


def _rollback_revision(
    group: SynchronizedCaptureGroup,
    revision_directory: Path,
    clock: Callable[[], datetime],
) -> None:
    active = group.measurement_directory / "yolo"
    previous = revision_directory / "evidence"
    if previous.is_dir():
        if active.exists():
            failed = revision_directory / "failed_reprocess_evidence"
            if failed.exists():
                raise SynchronizedProcessingError(f"Rollback target already exists: {failed}")
            if any(active.iterdir()):
                _archive_directory(active, failed)
        active.mkdir(parents=True, exist_ok=True)
        _copy_directory_contents_verified(previous, active)
    recovery = revision_directory / "recovery"
    restore_targets = {
        "yolo_measurements.xlsx": group.session_directory / "yolo_measurements.xlsx",
        "yolo_measurements.csv": group.session_directory / "yolo_measurements.csv",
        "session_manifest.json": group.session_directory / "session_manifest.json",
    }
    for name, target in restore_targets.items():
        backup = recovery / name
        if backup.is_file():
            _replace_from_copy(backup, target)
    config_backups = [
        path for path in recovery.glob("yolo_effective_config_*.yaml") if path.is_file()
    ]
    for backup in config_backups:
        _replace_from_copy(
            backup,
            group.session_directory / "provenance" / backup.name,
        )
    revision = _read_json(revision_directory / "revision.json")
    revision["state"] = "ROLLED_BACK_AFTER_ERROR"
    revision["rolled_back_utc"] = format_utc(clock())
    _atomic_json(revision_directory / "revision.json", revision)


def _archive_interrupted_evidence(
    measurement: Path,
    active: Path,
    revision: int,
    clock: Callable[[], datetime],
) -> None:
    directory = measurement / "yolo_revisions" / f"interrupted_{revision:04d}"
    directory.mkdir(parents=True, exist_ok=False)
    _archive_directory(active, directory / "evidence")
    _atomic_json(
        directory / "revision.json",
        {
            "revision_schema_version": 1,
            "state": "INTERRUPTED_EVIDENCE_ARCHIVED",
            "processing_revision": revision,
            "archived_utc": format_utc(clock()),
        },
    )


def _rename_directory_with_retry(
    source: Path,
    target: Path,
    *,
    attempts: int = 30,
    delay_seconds: float = 0.2,
) -> None:
    """Rename evidence despite short-lived Windows scanner/preview handles."""

    last_error: PermissionError | None = None
    for attempt in range(attempts):
        try:
            source.rename(target)
            return
        except PermissionError as exc:
            last_error = exc
            if attempt + 1 < attempts:
                time.sleep(delay_seconds)
    raise SynchronizedProcessingError(
        "Windows kept the YOLO evidence directory open, so it could not be "
        f"archived after {attempts} attempts: {source}. Close File Explorer/image "
        "preview windows and any program using this experiment folder, then rerun "
        "the same offline command. Existing evidence was not deleted."
    ) from last_error


def _archive_directory(source: Path, target: Path) -> None:
    """Archive a directory by rename, or by verified copy when its root is locked."""

    try:
        _rename_directory_with_retry(source, target)
        return
    except SynchronizedProcessingError:
        # Windows Explorer, antivirus, and indexers can hold a directory handle
        # while still permitting every child to be copied and removed. Commit a
        # verified snapshot before touching the active contents.
        _copy_directory_tree_verified(source, target)
        _clear_directory_contents(source)


def _copy_directory_tree_verified(source: Path, target: Path) -> None:
    if target.exists():
        raise SynchronizedProcessingError(f"Archive target already exists: {target}")
    temporary = target.with_name(f".{target.name}.{os.getpid()}.copying")
    if temporary.exists():
        shutil.rmtree(temporary)
    try:
        shutil.copytree(source, temporary, copy_function=shutil.copy2)
        _verify_directory_copy(source, temporary)
        temporary.rename(target)
    finally:
        if temporary.exists():
            shutil.rmtree(temporary)


def _copy_directory_contents_verified(source: Path, target: Path) -> None:
    if any(target.iterdir()):
        raise SynchronizedProcessingError(
            f"Rollback target directory is not empty: {target}"
        )
    shutil.copytree(source, target, dirs_exist_ok=True, copy_function=shutil.copy2)
    _verify_directory_copy(source, target)


def _verify_directory_copy(source: Path, target: Path) -> None:
    source_files = {
        path.relative_to(source).as_posix(): path
        for path in source.rglob("*")
        if path.is_file()
    }
    target_files = {
        path.relative_to(target).as_posix(): path
        for path in target.rglob("*")
        if path.is_file()
    }
    if source_files.keys() != target_files.keys():
        raise SynchronizedProcessingError(
            f"Evidence archive file list differs from its source: {target}"
        )
    for relative, source_file in source_files.items():
        target_file = target_files[relative]
        if (
            source_file.stat().st_size != target_file.stat().st_size
            or sha256_file(source_file) != sha256_file(target_file)
        ):
            raise SynchronizedProcessingError(
                f"Evidence archive verification failed for {relative}: {target}"
            )


def _clear_directory_contents(directory: Path) -> None:
    for child in tuple(directory.iterdir()):
        if child.is_dir() and not child.is_symlink():
            shutil.rmtree(child)
        else:
            child.unlink()
    if any(directory.iterdir()):
        raise SynchronizedProcessingError(
            f"Could not clear archived active evidence directory: {directory}"
        )


def _finalize_bookkeeping(
    group: SynchronizedCaptureGroup,
    record: CommonMeasurementRecord,
    profile: ModelProfile,
    revision: int,
    revision_directory: Path | None,
    clock: Callable[[], datetime],
) -> None:
    session = group.session_directory
    capture = _read_json(group.capture_manifest_path)
    capture["offline_yolo_processing"] = {
        "state": "COMPLETED",
        "processing_revision": revision,
        "profile": profile.name,
        "status": record.status,
        "valid": record.valid,
        "config_sha256": record.config_sha256,
        "model_sha256": record.calibration_or_model_sha256,
        "artifact_directory": record.artifact_directory,
        "completed_utc": format_utc(clock()),
    }
    _atomic_json(group.capture_manifest_path, capture)
    with _exclusive_lock(session / ".yolo_session.lock"):
        manifest_path = session / "session_manifest.json"
        manifest = _read_json(manifest_path)
        acquisitions = manifest.get("synchronized_acquisitions", [])
        if isinstance(acquisitions, list):
            for item in acquisitions:
                if isinstance(item, dict) and item.get("measurement_id") == group.measurement_id:
                    item["yolo_inference_deferred"] = False
                    item["yolo_inference_completed"] = True
                    item["yolo_processing_revision"] = revision
        if revision_directory is not None:
            revisions = manifest.setdefault("yolo_processing_revisions", [])
            revisions.append(
                {
                    "measurement_id": group.measurement_id,
                    "processing_revision": revision,
                    "profile": profile.name,
                    "previous_evidence": revision_directory.relative_to(session).as_posix(),
                    "previous_config_sha256": _read_json(
                        revision_directory / "revision.json"
                    ).get("previous_config_sha256"),
                    "previous_model_sha256": _read_json(
                        revision_directory / "revision.json"
                    ).get("previous_model_sha256"),
                    "new_config_sha256": record.config_sha256,
                    "new_model_sha256": record.calibration_or_model_sha256,
                    "completed_utc": format_utc(clock()),
                }
            )
            revision_data = _read_json(revision_directory / "revision.json")
            revision_data["state"] = "PRESERVED"
            revision_data["completed_utc"] = format_utc(clock())
            _atomic_json(revision_directory / "revision.json", revision_data)
        manifest["updated_utc"] = format_utc(clock())
        _atomic_json(manifest_path, manifest)


def _result_from_record(
    row: CommonMeasurementRecord,
    profile_name: str,
    action: str,
    revision: int,
    session: Path,
) -> SynchronizedGroupResult:
    artifact = _session_path(session, row.artifact_directory, "artifact_directory")
    details = _read_json(artifact / "result.json")
    return SynchronizedGroupResult(
        experiment_id=str(row.experiment_id),
        measurement_id=str(row.measurement_id),
        measurement_index=int(row.measurement_index),
        action=action,
        profile_name=profile_name,
        status=str(row.status),
        valid=row.valid,
        accepted_image_count=details.get("accepted_image_count"),
        rejected_image_count=details.get("rejected_image_count"),
        processing_revision=revision,
        artifact_directory=artifact,
    )


def _recorded_revision(journal_path: Path) -> int:
    if not journal_path.is_file():
        return 1
    value = _read_json(journal_path).get("processing_revision", 1)
    return int(value) if isinstance(value, int) and not isinstance(value, bool) else 1


def _next_revision(measurement: Path) -> int:
    journal = measurement / "yolo_processing_manifest.json"
    values = [_recorded_revision(journal)] if journal.is_file() else []
    root = measurement / "yolo_revisions"
    if root.is_dir():
        for child in root.iterdir():
            if not child.is_dir():
                continue
            tail = child.name.rsplit("_", 1)[-1]
            if tail.isdigit():
                values.append(int(tail))
    return max(values, default=0) + 1


def _write_journal(
    path: Path,
    group: SynchronizedCaptureGroup,
    **values: Any,
) -> None:
    data = {
        "processing_manifest_schema_version": 1,
        "experiment_id": group.experiment_id,
        "measurement_id": group.measurement_id,
        "measurement_index": group.measurement_index,
        "capture_manifest": group.capture_manifest_path.relative_to(
            group.session_directory
        ).as_posix(),
        **values,
    }
    _atomic_json(path, data)


def _depth_hashes(session: Path) -> dict[Path, str]:
    return {
        path: sha256_file(path)
        for path in (session / "depth_measurements.xlsx", session / "depth_measurements.csv")
        if path.is_file()
    }


def _assert_depth_unchanged(session: Path, before: Mapping[Path, str]) -> None:
    after = _depth_hashes(session)
    if dict(before) != after:
        raise SynchronizedProcessingError(
            "Depth workbook/CSV changed during offline YOLO processing; refusing success."
        )


def _session_path(session: Path, value: Any, field: str) -> Path:
    if not isinstance(value, str) or not value.strip():
        raise SynchronizedProcessingError(f"Completed YOLO row has blank {field}.")
    candidate = Path(value)
    if candidate.is_absolute():
        raise SynchronizedProcessingError(f"{field} must be session-relative.")
    path = (session / candidate).resolve()
    try:
        path.relative_to(session.resolve())
    except ValueError as exc:
        raise SynchronizedProcessingError(f"{field} escapes the session: {value}") from exc
    return path


def _artifact_child(root: Path, relative: str) -> Path:
    candidate = Path(relative)
    if candidate.is_absolute():
        raise SynchronizedProcessingError("Artifact manifest paths must be relative.")
    path = (root / candidate).resolve()
    try:
        path.relative_to(root.resolve())
    except ValueError as exc:
        raise SynchronizedProcessingError(
            f"Artifact manifest path escapes evidence root: {relative}"
        ) from exc
    return path


def _read_json(path: Path) -> dict[str, Any]:
    if not path.is_file():
        raise SynchronizedProcessingError(f"Required JSON file not found: {path}")
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise SynchronizedProcessingError(f"Could not read JSON {path}: {exc}") from exc
    if not isinstance(value, dict):
        raise SynchronizedProcessingError(f"JSON root must be an object: {path}")
    return value


def _required_text(value: Mapping[str, Any], key: str, source: str) -> str:
    item = value.get(key)
    if not isinstance(item, str) or not item.strip():
        raise SynchronizedProcessingError(f"{source} requires non-blank {key}.")
    return item.strip()


def _atomic_json(path: Path, value: Any) -> None:
    _atomic_bytes(
        path,
        (json.dumps(value, indent=2, ensure_ascii=False, sort_keys=True) + "\n").encode(
            "utf-8"
        ),
    )


def _atomic_bytes(path: Path, payload: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, raw = tempfile.mkstemp(
        prefix=f".{path.stem}.", suffix=".tmp", dir=path.parent
    )
    temporary = Path(raw)
    try:
        with os.fdopen(descriptor, "wb") as stream:
            stream.write(payload)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def _copy_file(source: Path, target: Path) -> None:
    if target.exists():
        raise FileExistsError(f"Refusing to overwrite recovery evidence: {target}")
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_name(f".{target.name}.{os.getpid()}.tmp")
    try:
        shutil.copyfile(source, temporary)
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)


def _replace_from_copy(source: Path, target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    descriptor, raw = tempfile.mkstemp(
        prefix=f".{target.stem}.", suffix=".restore", dir=target.parent
    )
    os.close(descriptor)
    temporary = Path(raw)
    try:
        shutil.copyfile(source, temporary)
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)


@contextmanager
def _exclusive_lock(path: Path, timeout_seconds: float = 10.0) -> Iterator[None]:
    path.parent.mkdir(parents=True, exist_ok=True)
    deadline = time.monotonic() + timeout_seconds
    descriptor: int | None = None
    while descriptor is None:
        try:
            descriptor = os.open(path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
        except FileExistsError as exc:
            if time.monotonic() >= deadline:
                raise TimeoutError(f"Timed out waiting for offline YOLO lock: {path}") from exc
            time.sleep(0.05)
    try:
        os.write(descriptor, f"pid={os.getpid()}\n".encode("ascii"))
        os.close(descriptor)
        descriptor = None
        yield
    finally:
        if descriptor is not None:
            os.close(descriptor)
        path.unlink(missing_ok=True)
