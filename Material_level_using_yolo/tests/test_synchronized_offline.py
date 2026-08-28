from __future__ import annotations

from dataclasses import replace
from datetime import datetime, timedelta, timezone
import json
from pathlib import Path
from types import SimpleNamespace

import numpy as np
from openpyxl import load_workbook
import pytest

from experiment_records import (
    COMMON_COLUMNS,
    CommonMeasurementRecord,
    MeasurementWorkbookStore,
    format_utc,
    make_measurement_id,
    prepare_record,
    sha256_file,
    with_estimate,
)
from material_level_yolo.config import load_config
from material_level_yolo.domain import InferenceOutput, ProjectConfig, SemanticClassMap
from material_level_yolo.errors import ConfigurationError
from material_level_yolo.image_io import write_image
from material_level_yolo.synchronized import (
    SynchronizedProcessingError,
    _groups_from_index,
    _rename_directory_with_retry,
    discover_synchronized_groups,
    process_synchronized_captures,
)
from material_level_yolo.synchronized_cli import parse_args as parse_synchronized_args


PROJECT_ROOT = Path(__file__).resolve().parents[1]
TRIGGER = datetime(2026, 8, 14, 10, 0, 0, tzinfo=timezone.utc)
ZERO_HASH = "0" * 64
ONE_HASH = "1" * 64


def test_resume_index_selects_only_remaining_ordered_groups() -> None:
    groups = tuple(SimpleNamespace(measurement_index=index) for index in range(1, 25))
    selected = _groups_from_index(groups, 18)
    assert [group.measurement_index for group in selected] == list(range(18, 25))
    with pytest.raises(SynchronizedProcessingError, match="No capture groups"):
        _groups_from_index(groups, 25)
    with pytest.raises(SynchronizedProcessingError, match="at least 1"):
        _groups_from_index(groups, 0)


def test_synchronized_cli_accepts_resume_index() -> None:
    args = parse_synchronized_args(["session", "--start-measurement-index", "18"])
    assert args.start_measurement_index == 18


def test_windows_directory_rename_retries_transient_permission_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "yolo"
    target = tmp_path / "revision" / "evidence"
    source.mkdir()
    target.parent.mkdir()
    (source / "result.json").write_text("{}", encoding="utf-8")
    original_rename = Path.rename
    calls = 0

    def flaky_rename(path: Path, destination: Path):
        nonlocal calls
        if path == source and calls < 2:
            calls += 1
            raise PermissionError(5, "simulated transient Windows lock", str(path))
        return original_rename(path, destination)

    monkeypatch.setattr(Path, "rename", flaky_rename)
    monkeypatch.setattr("material_level_yolo.synchronized.time.sleep", lambda _delay: None)
    _rename_directory_with_retry(source, target, attempts=3, delay_seconds=0.0)
    assert calls == 2
    assert (target / "result.json").is_file()


class FakeBoxes:
    def __init__(self, classes, confidences) -> None:
        self.cls = np.asarray(classes, dtype=np.float32)
        self.conf = np.asarray(confidences, dtype=np.float32)
        self.xyxy = np.empty((0, 4), dtype=np.float32)


class FakeMasks:
    def __init__(self, masks) -> None:
        self.data = np.asarray(masks, dtype=np.float32)


class FakeResult:
    def __init__(self, classes, confidences, masks) -> None:
        self.boxes = FakeBoxes(classes, confidences)
        self.masks = FakeMasks(masks)


class FakeAdapter:
    def __init__(self, results, *, model_hash: str = "a" * 64, fail_predict=False) -> None:
        self.results = list(results)
        self.mapping = SemanticClassMap(
            material_id=7,
            material_name="Powder",
            empty_id=3,
            empty_name="Empty",
            model_names={3: "Empty", 7: "Powder"},
        )
        self.weights_sha256 = model_hash
        self.predict_calls = 0
        self.fail_predict = fail_predict

    def load(self):
        return self.mapping

    def predict(self, image):
        del image
        if self.fail_predict:
            raise RuntimeError("simulated interrupted inference")
        result = self.results[self.predict_calls]
        self.predict_calls += 1
        return InferenceOutput(
            raw_results=(result,),
            class_map=self.mapping,
            model_task="segment",
            weights_sha256=self.weights_sha256,
        )


def semantic_result(percent: float, *, ambiguous: bool = False) -> FakeResult:
    height, width = 100, 80
    boundary = round(height * (1.0 - percent / 100.0))
    material = np.zeros((height, width), dtype=np.float32)
    empty = np.zeros_like(material)
    material[boundary:] = 1.0
    empty[:boundary] = 1.0
    if ambiguous:
        # Both roles are detected, but the masks have the physically reversed
        # orientation and must be retained as an invalid image diagnostic.
        return FakeResult([7, 3], [0.95, 0.96], [empty, material])
    return FakeResult([7, 3], [0.95, 0.96], [material, empty])


def frozen_project(tmp_path: Path) -> ProjectConfig:
    project = load_config(PROJECT_ROOT / "config.yaml")
    reference = tmp_path / "calibration" / "tube_roi.json"
    reference.parent.mkdir(parents=True, exist_ok=True)
    reference.write_text("{}\n", encoding="utf-8")
    profiles = {
        name: replace(
            profile,
            tube_roi=replace(
                profile.tube_roi,
                frozen=True,
                left=0.05,
                right=0.95,
                calibration_reference=reference,
            ),
        )
        for name, profile in project.profiles.items()
    }
    return replace(project, profiles=profiles)


def make_session(
    tmp_path: Path,
    *,
    material_name: str = "Powder",
    experiment_id: str = "offline-synchronized-study",
) -> tuple[Path, str]:
    session = tmp_path / "Résultats synchronisés avec espaces" / experiment_id
    measurement_id = make_measurement_id(1, TRIGGER)
    measurement = session / "measurements" / measurement_id
    raw = measurement / "c920" / "raw"
    raw.mkdir(parents=True)
    frames = []
    for index in range(1, 7):
        actual = TRIGGER + timedelta(seconds=(index - 1) * 0.2)
        path = write_image(
            raw / f"c920_{index:06d}_échantillon.png",
            np.full((100, 80, 3), index * 20, dtype=np.uint8),
        )
        frames.append(
            {
                "index": index,
                "path": path.relative_to(session).as_posix(),
                "sha256": sha256_file(path),
                "target_offset_seconds": (index - 1) * 0.2,
                "actual_utc": format_utc(actual),
            }
        )
    capture_manifest = measurement / "capture_manifest.json"
    capture_manifest.write_text(
        json.dumps(
            {
                "capture_manifest_schema_version": 1,
                "experiment_id": experiment_id,
                "measurement_id": measurement_id,
                "measurement_index": 1,
                "state": "SAVED",
                "trigger_time_utc": format_utc(TRIGGER),
                "c920": {"planned_count": 6, "frames": frames},
                "depth": {"status": "captured"},
                "yolo_loaded_or_run": False,
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    (session / "session_manifest.json").write_text(
        json.dumps(
            {
                "schema_version": 1,
                "experiment_id": experiment_id,
                "purpose": "offline synchronized test",
                "measurements": [
                    {
                        "measurement_id": measurement_id,
                        "measurement_index": 1,
                        "method": "depth",
                        "status": "complete_valid",
                    },
                    {
                        "measurement_id": measurement_id,
                        "measurement_index": 1,
                        "method": "yolo",
                        "status": "pending_offline_inference",
                    },
                ],
                "synchronized_acquisitions": [
                    {
                        "measurement_id": measurement_id,
                        "yolo_inference_deferred": True,
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    yolo = prepare_record(
        CommonMeasurementRecord(
            experiment_id=experiment_id,
            measurement_id=measurement_id,
            measurement_index=1,
            method="yolo",
            acquisition_mode="synchronized",
            trigger_time_utc=TRIGGER,
            capture_start_utc=TRIGGER,
            capture_end_utc=TRIGGER + timedelta(seconds=1),
            material_name=material_name,
            total_capacity_ml=200.0,
            bulk_density_g_per_ml=0.5,
            total_possible_weight_g=100.0,
            reference_material_weight_g=25.0,
            valid=None,
            status="pending_offline_inference",
            notes="operator metadata retained",
            artifact_directory=f"measurements/{measurement_id}/c920",
            source_artifact=f"measurements/{measurement_id}/capture_manifest.json",
        )
    )
    MeasurementWorkbookStore(session / "yolo_measurements.xlsx", method="yolo").upsert(yolo)
    depth = prepare_record(
        with_estimate(
            CommonMeasurementRecord(
                experiment_id=experiment_id,
                measurement_id=measurement_id,
                measurement_index=1,
                method="depth",
                acquisition_mode="synchronized",
                trigger_time_utc=TRIGGER,
                capture_start_utc=TRIGGER,
                capture_end_utc=TRIGGER + timedelta(seconds=0.8),
                processing_time_utc=TRIGGER + timedelta(seconds=1.2),
                material_name=material_name,
                total_capacity_ml=200.0,
                bulk_density_g_per_ml=0.5,
                total_possible_weight_g=100.0,
                reference_material_weight_g=25.0,
                valid=True,
                status="complete_valid",
                artifact_directory=f"measurements/{measurement_id}/depth",
                source_artifact=f"measurements/{measurement_id}/depth/result.json",
                config_sha256=ZERO_HASH,
                calibration_or_model_sha256=ONE_HASH,
            ),
            42.0,
        )
    )
    MeasurementWorkbookStore(session / "depth_measurements.xlsx", method="depth").upsert(depth)
    return session, measurement_id


def adapter_factory(results, *, model_hash="a" * 64, fail_predict=False):
    created = []

    def factory(profile):
        assert profile.name == "powder"
        adapter = FakeAdapter(results, model_hash=model_hash, fail_predict=fail_predict)
        created.append(adapter)
        return adapter

    return factory, created


def test_manifest_discovery_processes_pending_row_and_second_run_is_idempotent(
    tmp_path: Path,
) -> None:
    session, measurement_id = make_session(tmp_path)
    groups = discover_synchronized_groups(session)
    assert len(groups) == 1
    assert groups[0].measurement_id == measurement_id
    assert discover_synchronized_groups(groups[0].measurement_directory) == groups
    depth_before = {
        name: (session / name).read_bytes()
        for name in ("depth_measurements.xlsx", "depth_measurements.csv")
    }
    factory, created = adapter_factory([semantic_result(50.0) for _ in range(6)])
    report = process_synchronized_captures(
        frozen_project(tmp_path),
        session,
        adapter_factory=factory,
        now_utc=lambda: TRIGGER + timedelta(seconds=2),
    )
    assert report.groups[0].action == "processed"
    assert report.groups[0].measurement_id == measurement_id
    assert report.groups[0].profile_name == "powder"
    assert created[0].predict_calls == 6
    row = MeasurementWorkbookStore(
        session / "yolo_measurements.xlsx", method="yolo"
    ).load_records()[0]
    assert row.measurement_id == measurement_id
    assert row.status == "complete_valid"
    assert row.estimated_material_percent == pytest.approx(50.0)
    assert row.source_artifact == f"measurements/{measurement_id}/capture_manifest.json"
    for name, payload in depth_before.items():
        assert (session / name).read_bytes() == payload
    depth_book = load_workbook(session / "depth_measurements.xlsx", read_only=True)
    yolo_book = load_workbook(session / "yolo_measurements.xlsx", read_only=True)
    try:
        assert tuple(cell.value for cell in depth_book["measurements"][1]) == COMMON_COLUMNS
        assert tuple(cell.value for cell in yolo_book["measurements"][1]) == COMMON_COLUMNS
    finally:
        depth_book.close()
        yolo_book.close()

    def forbidden_factory(_profile):
        raise AssertionError("idempotent verification must not load YOLO")

    second = process_synchronized_captures(
        frozen_project(tmp_path),
        session,
        adapter_factory=forbidden_factory,
    )
    assert second.groups[0].action == "verified_existing"
    assert len(
        MeasurementWorkbookStore(
            session / "yolo_measurements.xlsx", method="yolo"
        ).load_records()
    ) == 1


def test_force_reprocess_preserves_revision_when_model_and_config_hash_change(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    session, measurement_id = make_session(tmp_path)
    project = frozen_project(tmp_path)
    first_factory, _ = adapter_factory([semantic_result(45.0) for _ in range(6)])
    process_synchronized_captures(
        project,
        session,
        adapter_factory=first_factory,
        now_utc=lambda: TRIGGER + timedelta(seconds=2),
    )
    first = MeasurementWorkbookStore(
        session / "yolo_measurements.xlsx", method="yolo"
    ).load_records()[0]
    changed_profile = replace(
        project.select_profile("powder"),
        aggregation=replace(
            project.select_profile("powder").aggregation,
            maximum_spread_percentage_points=7.5,
        ),
    )
    changed_project = ProjectConfig(
        schema_version=project.schema_version,
        config_path=project.config_path,
        default_profile=project.default_profile,
        profiles={**project.profiles, "powder": changed_profile},
    )
    second_factory, _ = adapter_factory(
        [semantic_result(55.0) for _ in range(6)], model_hash="b" * 64
    )
    active_evidence = session / "measurements" / measurement_id / "yolo"
    original_rename = Path.rename

    def persistently_locked_root(path: Path, destination: Path):
        if path == active_evidence:
            raise PermissionError(5, "simulated persistent Windows directory lock", str(path))
        return original_rename(path, destination)

    monkeypatch.setattr(Path, "rename", persistently_locked_root)
    monkeypatch.setattr("material_level_yolo.synchronized.time.sleep", lambda _delay: None)
    report = process_synchronized_captures(
        changed_project,
        session,
        force_reprocess=True,
        adapter_factory=second_factory,
        now_utc=lambda: TRIGGER + timedelta(seconds=3),
    )
    assert report.groups[0].action == "reprocessed_revision"
    current = MeasurementWorkbookStore(
        session / "yolo_measurements.xlsx", method="yolo"
    ).load_records()[0]
    assert current.measurement_id == measurement_id
    assert current.calibration_or_model_sha256 == "b" * 64
    assert current.config_sha256 != first.config_sha256
    revision = session / "measurements" / measurement_id / "yolo_revisions" / "revision_0002"
    assert (revision / "evidence" / "artifact_manifest.json").is_file()
    previous = json.loads((revision / "previous_record.json").read_text(encoding="utf-8"))
    assert previous["calibration_or_model_sha256"] == "a" * 64
    assert (active_evidence / "artifact_manifest.json").is_file()
    session_manifest = json.loads((session / "session_manifest.json").read_text(encoding="utf-8"))
    assert session_manifest["yolo_processing_revisions"][-1]["measurement_id"] == measurement_id


@pytest.mark.parametrize("failure", ["missing", "corrupt"])
def test_missing_or_corrupted_source_fails_before_model_and_keeps_pending_row(
    tmp_path: Path,
    failure: str,
) -> None:
    session, _measurement_id = make_session(tmp_path, experiment_id=f"integrity-{failure}")
    source = discover_synchronized_groups(session)[0].source_paths[2]
    if failure == "missing":
        source.unlink()
    else:
        source.write_bytes(b"corrupted source bytes")

    def forbidden_factory(_profile):
        raise AssertionError("model must not load before source verification")

    with pytest.raises(SynchronizedProcessingError, match="missing|hash mismatch"):
        process_synchronized_captures(
            frozen_project(tmp_path),
            session,
            adapter_factory=forbidden_factory,
        )
    row = MeasurementWorkbookStore(
        session / "yolo_measurements.xlsx", method="yolo"
    ).load_records()[0]
    assert row.status == "pending_offline_inference"
    assert row.estimated_material_percent is None


def test_invalid_and_partially_valid_frames_retain_all_diagnostic_reasons(
    tmp_path: Path,
) -> None:
    session, measurement_id = make_session(tmp_path, experiment_id="partially-valid")
    results = [semantic_result(50.0) for _ in range(3)] + [
        semantic_result(50.0, ambiguous=True) for _ in range(3)
    ]
    factory, _ = adapter_factory(results)
    report = process_synchronized_captures(
        frozen_project(tmp_path),
        session,
        adapter_factory=factory,
        now_utc=lambda: TRIGGER + timedelta(seconds=2),
    )
    assert report.groups[0].valid is True
    assert report.groups[0].accepted_image_count == 3
    assert report.groups[0].rejected_image_count == 3
    workbook = load_workbook(session / "yolo_measurements.xlsx", data_only=True)
    try:
        details = workbook["yolo_image_details"]
        headers = [cell.value for cell in details[1]]
        reason_column = headers.index("rejection_reasons") + 1
        reasons = [details.cell(row, reason_column).value for row in range(2, 8)]
        assert sum("implausible_material_empty_orientation" in (value or "") for value in reasons) == 3
    finally:
        workbook.close()
    assert all(
        row.measurement_id == measurement_id
        for row in MeasurementWorkbookStore(
            session / "yolo_measurements.xlsx", method="yolo"
        ).load_records()
    )

    invalid_session, _ = make_session(tmp_path, experiment_id="below-minimum")
    invalid_results = [semantic_result(50.0, ambiguous=True) for _ in range(6)]
    invalid_factory, _ = adapter_factory(invalid_results)
    invalid_report = process_synchronized_captures(
        frozen_project(tmp_path),
        invalid_session,
        adapter_factory=invalid_factory,
        now_utc=lambda: TRIGGER + timedelta(seconds=2),
    )
    invalid_row = MeasurementWorkbookStore(
        invalid_session / "yolo_measurements.xlsx", method="yolo"
    ).load_records()[0]
    assert invalid_report.groups[0].valid is False
    assert invalid_row.status == "complete_invalid"
    assert invalid_row.estimated_material_percent is None
    assert invalid_row.estimated_material_volume_ml is None


def test_interrupted_pending_attempt_is_archived_and_retry_completes(tmp_path: Path) -> None:
    session, measurement_id = make_session(tmp_path, experiment_id="retry-after-interruption")
    failing_factory, _ = adapter_factory(
        [semantic_result(50.0) for _ in range(6)], fail_predict=True
    )
    with pytest.raises(RuntimeError, match="interrupted inference"):
        process_synchronized_captures(
            frozen_project(tmp_path),
            session,
            adapter_factory=failing_factory,
            now_utc=lambda: TRIGGER + timedelta(seconds=2),
        )
    pending = MeasurementWorkbookStore(
        session / "yolo_measurements.xlsx", method="yolo"
    ).load_records()[0]
    assert pending.status == "pending_offline_inference"
    journal = session / "measurements" / measurement_id / "yolo_processing_manifest.json"
    assert json.loads(journal.read_text(encoding="utf-8"))["state"] == "ERROR"

    retry_factory, _ = adapter_factory([semantic_result(50.0) for _ in range(6)])
    report = process_synchronized_captures(
        frozen_project(tmp_path),
        session,
        adapter_factory=retry_factory,
        now_utc=lambda: TRIGGER + timedelta(seconds=3),
    )
    assert report.groups[0].action == "processed"
    assert list(
        (session / "measurements" / measurement_id / "yolo_revisions").glob(
            "interrupted_*/evidence"
        )
    )
    completed = MeasurementWorkbookStore(
        session / "yolo_measurements.xlsx", method="yolo"
    ).load_records()[0]
    assert completed.status == "complete_valid"


def test_failed_forced_revision_rolls_back_previous_completed_result(tmp_path: Path) -> None:
    session, measurement_id = make_session(tmp_path, experiment_id="force-rollback")
    project = frozen_project(tmp_path)
    first_factory, _ = adapter_factory([semantic_result(47.0) for _ in range(6)])
    process_synchronized_captures(
        project,
        session,
        adapter_factory=first_factory,
        now_utc=lambda: TRIGGER + timedelta(seconds=2),
    )
    store = MeasurementWorkbookStore(session / "yolo_measurements.xlsx", method="yolo")
    previous = store.load_records()[0]
    previous_manifest_hash = sha256_file(
        session / "measurements" / measurement_id / "yolo" / "artifact_manifest.json"
    )
    failing_factory, _ = adapter_factory(
        [semantic_result(60.0) for _ in range(6)],
        model_hash="b" * 64,
        fail_predict=True,
    )
    with pytest.raises(RuntimeError, match="interrupted inference"):
        process_synchronized_captures(
            project,
            session,
            force_reprocess=True,
            adapter_factory=failing_factory,
            now_utc=lambda: TRIGGER + timedelta(seconds=3),
        )
    restored = store.load_records()[0]
    assert restored.to_row() == previous.to_row()
    assert sha256_file(
        session / "measurements" / measurement_id / "yolo" / "artifact_manifest.json"
    ) == previous_manifest_hash
    revision = session / "measurements" / measurement_id / "yolo_revisions" / "revision_0002"
    assert json.loads((revision / "revision.json").read_text(encoding="utf-8"))["state"] == (
        "ROLLED_BACK_AFTER_ERROR"
    )

    def forbidden_factory(_profile):
        raise AssertionError("rolled-back completed result should verify without inference")

    report = process_synchronized_captures(
        project,
        session,
        adapter_factory=forbidden_factory,
    )
    assert report.groups[0].action == "verified_existing"


def test_recorded_material_requires_unambiguous_configured_profile(tmp_path: Path) -> None:
    session, _ = make_session(tmp_path, material_name="unconfigured resin")
    with pytest.raises(ConfigurationError, match="configured for recorded material"):
        process_synchronized_captures(
            frozen_project(tmp_path),
            session,
            adapter_factory=lambda _profile: pytest.fail("must not load model"),
        )


def test_explicit_profile_recovers_blank_recorded_material_without_inventing_it(
    tmp_path: Path,
) -> None:
    session, measurement_id = make_session(
        tmp_path,
        material_name=None,
        experiment_id="legacy-blank-material",
    )
    factory, created = adapter_factory([semantic_result(50.0) for _ in range(6)])

    report = process_synchronized_captures(
        frozen_project(tmp_path),
        session,
        profile_name="powder",
        adapter_factory=factory,
        now_utc=lambda: TRIGGER + timedelta(seconds=2),
    )

    assert report.groups[0].action == "processed"
    assert report.groups[0].measurement_id == measurement_id
    assert report.groups[0].profile_name == "powder"
    assert created[0].predict_calls == 6
    row = MeasurementWorkbookStore(
        session / "yolo_measurements.xlsx", method="yolo"
    ).load_records()[0]
    assert row.status == "complete_valid"
    assert row.material_name is None
