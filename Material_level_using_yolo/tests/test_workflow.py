from __future__ import annotations

from dataclasses import replace
from datetime import datetime, timedelta, timezone
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
import pytest

from experiment_records import (
    COMMON_COLUMNS,
    CommonMeasurementRecord,
    MeasurementWorkbookStore,
    format_utc,
    make_measurement_id,
    prepare_record,
    sha256_file,
)
from material_level_yolo.config import load_config
from material_level_yolo.domain import InferenceOutput, SemanticClassMap
from material_level_yolo.image_io import write_image
from material_level_yolo.roi_policy import apply_roi_mode
from material_level_yolo.workflow import (
    YOLO_IMAGE_DETAIL_COLUMNS,
    OfflineMeasurementRequest,
    process_saved_images,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]
TRIGGER = datetime(2026, 8, 13, 15, 0, 0, 123456, tzinfo=timezone.utc)


class FakeBoxes:
    def __init__(self, classes, confidences, xyxy=None) -> None:
        self.cls = np.asarray(classes, dtype=np.float32)
        self.conf = np.asarray(confidences, dtype=np.float32)
        self.xyxy = np.asarray(xyxy if xyxy is not None else [], dtype=np.float32)


class FakeMasks:
    def __init__(self, masks) -> None:
        self.data = np.asarray(masks, dtype=np.float32)


class FakeResult:
    def __init__(self, classes, confidences, masks, xyxy=None) -> None:
        self.boxes = FakeBoxes(classes, confidences, xyxy)
        self.masks = FakeMasks(masks)


class FakeAdapter:
    def __init__(self, results) -> None:
        self.results = list(results)
        self.mapping = SemanticClassMap(
            material_id=7,
            material_name="Powder",
            empty_id=3,
            empty_name="Empty",
            model_names={3: "Empty", 7: "Powder"},
        )
        self.weights_sha256 = "a" * 64
        self.resolved_device = "cpu"
        self.predict_calls = 0

    def load(self):
        return self.mapping

    def predict(self, image):
        result = self.results[self.predict_calls]
        self.predict_calls += 1
        return InferenceOutput(
            raw_results=(result,),
            class_map=self.mapping,
            model_task="segment",
            weights_sha256=self.weights_sha256,
            inference_device=self.resolved_device,
        )


def semantic_result(percent: float, *, duplicate_material: bool = False) -> FakeResult:
    height, width = 100, 80
    boundary = round(height * (1.0 - percent / 100.0))
    material = np.zeros((height, width), dtype=np.float32)
    empty = np.zeros_like(material)
    material[boundary:] = 1.0
    empty[:boundary] = 1.0
    if duplicate_material:
        return FakeResult(
            [7, 7, 3],
            [0.95, 0.90, 0.96],
            [material, material, empty],
        )
    return FakeResult([7, 3], [0.95, 0.96], [material, empty])


def material_only_result(percent: float = 40.0) -> FakeResult:
    height, width = 100, 80
    boundary = round(height * (1.0 - percent / 100.0))
    material = np.zeros((height, width), dtype=np.float32)
    material[boundary:] = 1.0
    return FakeResult([7], [0.95], [material])


def source_images(tmp_path: Path, count: int) -> tuple[Path, ...]:
    folder = tmp_path / "Images d'étude avec espaces"
    values = []
    for index in range(1, count + 1):
        image = np.full((100, 80, 3), 20 * index, dtype=np.uint8)
        values.append(write_image(folder / f"échantillon_{index:02d}.png", image))
    return tuple(values)


def frozen_project(tmp_path: Path):
    project = load_config(PROJECT_ROOT / "config.yaml")
    reference = tmp_path / "calibration" / "tube_roi.json"
    reference.parent.mkdir(parents=True, exist_ok=True)
    reference.write_text("{}\n", encoding="utf-8")
    profiles = {
        name: replace(
            profile,
            tube_roi=replace(
                profile.tube_roi,
                frozen=True,
                left=0.05,
                right=0.95,
                calibration_reference=reference,
            ),
        )
        for name, profile in project.profiles.items()
    }
    return replace(project, profiles=profiles)


def test_mocked_adapter_group_writes_common_row_details_and_evidence(tmp_path: Path) -> None:
    project = frozen_project(tmp_path)
    sources = source_images(tmp_path, 6)
    original_hashes = {path: sha256_file(path) for path in sources}
    adapter = FakeAdapter(
        [semantic_result(value) for value in (49.0, 50.0, 51.0, 50.0, 49.0, 51.0)]
    )
    outcome = process_saved_images(
        project,
        "powder",
        sources,
        OfflineMeasurementRequest(
            experiment_id="yolo-mock-study",
            total_capacity_ml=200.0,
            output_root=(tmp_path / "Résultats YOLO").resolve(),
            material_name="operator supplied powder",
            trigger_time_utc=TRIGGER,
            software_repository=PROJECT_ROOT.parent,
        ),
        adapter=adapter,
        now_utc=lambda: TRIGGER,
    )
    assert adapter.predict_calls == 6
    assert outcome.group.valid
    assert outcome.record.method == "yolo"
    assert outcome.record.estimated_material_percent == pytest.approx(50.0)
    assert outcome.record.estimated_material_volume_ml == pytest.approx(100.0)
    assert outcome.record.estimated_empty_percent == pytest.approx(50.0)
    assert outcome.record.estimated_empty_volume_ml == pytest.approx(100.0)
    assert all(sha256_file(path) == digest for path, digest in original_hashes.items())

    workbook_path = outcome.session_directory / "yolo_measurements.xlsx"
    store = MeasurementWorkbookStore(workbook_path, method="yolo")
    assert store.verify_mirror()
    workbook = load_workbook(workbook_path, data_only=True)
    assert tuple(cell.value for cell in workbook["measurements"][1]) == COMMON_COLUMNS
    details = workbook["yolo_image_details"]
    assert tuple(cell.value for cell in details[1]) == YOLO_IMAGE_DETAIL_COLUMNS
    assert details.max_row == 7
    assert all(details.cell(row, YOLO_IMAGE_DETAIL_COLUMNS.index("valid") + 1).value for row in range(2, 8))

    artifact_files = {
        path.relative_to(outcome.artifact_directory).as_posix()
        for path in outcome.artifact_directory.rglob("*")
        if path.is_file()
    }
    assert "result.json" in artifact_files
    assert "image_details.csv" in artifact_files
    assert "artifact_manifest.json" in artifact_files
    result_payload = json.loads(outcome.result_path.read_text(encoding="utf-8"))
    assert result_payload["requested_inference_device"] == "auto"
    assert result_payload["resolved_inference_device"] == "cpu"
    assert len([item for item in artifact_files if item.startswith("source_images/")]) == 6
    assert len([item for item in artifact_files if item.startswith("overlays/")]) == 6
    assert len([item for item in artifact_files if item.startswith("masks/")]) == 12
    manifest = json.loads(outcome.artifact_manifest_path.read_text(encoding="utf-8"))
    assert all(
        sha256_file(outcome.artifact_directory / item["path"]) == item["sha256"]
        for item in manifest["artifacts"]
    )


def test_unanimous_material_only_group_saves_overlay_without_image_candidate(
    tmp_path: Path,
) -> None:
    project = apply_roi_mode(load_config(PROJECT_ROOT / "config.yaml"), "whole_image")
    sources = source_images(tmp_path, 6)
    outcome = process_saved_images(
        project,
        "powder",
        sources,
        OfflineMeasurementRequest(
            experiment_id="single-class-overlay",
            total_capacity_ml=200.0,
            output_root=(tmp_path / "single class results").resolve(),
            material_name="white powder",
            trigger_time_utc=TRIGGER,
            software_repository=PROJECT_ROOT.parent,
        ),
        adapter=FakeAdapter([material_only_result() for _ in sources]),
        now_utc=lambda: TRIGGER,
    )
    assert outcome.group.valid
    assert outcome.group.material_percent == 100.0
    assert outcome.record.estimated_material_volume_ml == 200.0
    assert all(image.candidate_material_percent is None for image in outcome.group.images)
    assert len(list((outcome.artifact_directory / "overlays").glob("*.png"))) == 6


def test_multiple_model_regions_use_highest_confidence_and_remain_diagnostic(
    tmp_path: Path,
) -> None:
    project = frozen_project(tmp_path)
    sources = source_images(tmp_path, 3)
    outcome = process_saved_images(
        project,
        "powder",
        sources,
        OfflineMeasurementRequest(
            experiment_id="ambiguous-regions",
            total_capacity_ml=150.0,
            output_root=(tmp_path / "Sortie ambiguë").resolve(),
            trigger_time_utc=TRIGGER,
        ),
        adapter=FakeAdapter([semantic_result(50.0, duplicate_material=True) for _ in sources]),
        now_utc=lambda: TRIGGER,
    )
    assert outcome.group.valid
    assert outcome.record.status == "complete_valid"
    assert outcome.record.valid is True
    assert outcome.record.estimated_material_percent == pytest.approx(50.0)
    workbook = load_workbook(outcome.session_directory / "yolo_measurements.xlsx", data_only=True)
    count_column = YOLO_IMAGE_DETAIL_COLUMNS.index("material_instance_count") + 1
    assert all(
        workbook["yolo_image_details"].cell(row, count_column).value == 2
        for row in range(2, 5)
    )


def test_explicit_synchronized_resume_replaces_pending_row_with_same_identity(
    tmp_path: Path,
) -> None:
    project = frozen_project(tmp_path)
    output_root = (tmp_path / "Résultats synchronisés").resolve()
    experiment_id = "synchronized-offline-resume"
    measurement_id = make_measurement_id(1, TRIGGER)
    session = output_root / experiment_id
    (session / "measurements" / measurement_id / "c920" / "raw").mkdir(parents=True)
    pending = prepare_record(
        CommonMeasurementRecord(
            experiment_id=experiment_id,
            measurement_id=measurement_id,
            measurement_index=1,
            method="yolo",
            acquisition_mode="synchronized",
            trigger_time_utc=TRIGGER,
            capture_start_utc=TRIGGER,
            capture_end_utc=TRIGGER + timedelta(seconds=1),
            total_capacity_ml=200.0,
            valid=None,
            status="pending_offline_inference",
            artifact_directory=f"measurements/{measurement_id}/c920",
            source_artifact=f"measurements/{measurement_id}/capture_manifest.json",
        )
    )
    store = MeasurementWorkbookStore(session / "yolo_measurements.xlsx", method="yolo")
    store.upsert(pending)
    sources = source_images(tmp_path, 3)
    outcome = process_saved_images(
        project,
        "powder",
        sources,
        OfflineMeasurementRequest(
            experiment_id=experiment_id,
            total_capacity_ml=200.0,
            purpose="synchronized offline handoff",
            output_root=output_root,
            measurement_index=1,
            measurement_id=measurement_id,
            trigger_time_utc=TRIGGER,
            acquisition_mode="synchronized",
            capture_start_utc=TRIGGER,
            capture_end_utc=TRIGGER + timedelta(seconds=1),
            resume_pending_synchronized=True,
        ),
        adapter=FakeAdapter([semantic_result(50.0) for _ in sources]),
        now_utc=lambda: TRIGGER + timedelta(seconds=2),
    )
    assert outcome.record.measurement_id == measurement_id
    assert outcome.record.status == "complete_valid"
    records = store.load_records()
    assert len(records) == 1
    assert records[0].measurement_id == measurement_id
    assert records[0].status == "complete_valid"
    assert format_utc(records[0].capture_start_utc) == format_utc(TRIGGER)
