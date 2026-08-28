from __future__ import annotations

from dataclasses import replace
from datetime import datetime, timedelta, timezone
import json
from pathlib import Path
from types import SimpleNamespace
import subprocess
import sys
import threading

import numpy as np
from openpyxl import load_workbook
import pytest

from experiment_records import (
    COMMON_COLUMNS,
    CommonMeasurementRecord,
    MeasurementWorkbookStore,
    format_utc,
    make_measurement_id,
    next_measurement_identity,
    prepare_record,
    sha256_file,
    with_estimate,
)
from material_level_yolo.config import load_config as load_yolo_config
from material_level_yolo.domain import (
    CameraSettings,
    InferenceOutput,
    SemanticClassMap,
)
from material_level_yolo.synchronized import process_synchronized_captures
from run_experiment.errors import BusyTriggerError
from run_experiment.cli import _guided_measurement
from run_experiment.models import (
    AcquisitionSettings,
    ExperimentConfig,
    ExperimentInputs,
    ExperimentState,
    MeasurementReferenceInputs,
)
from run_experiment.core import SynchronizedExperiment


PROJECT_ROOT = Path(__file__).resolve().parents[1]
REPOSITORY_ROOT = PROJECT_ROOT.parent
TRIGGER = datetime(2026, 8, 14, 8, 0, 0, tzinfo=timezone.utc)
ZERO_HASH = "0" * 64
ONE_HASH = "1" * 64


class VirtualClock:
    def __init__(self) -> None:
        self.value = 100.0
        self.lock = threading.Lock()

    def monotonic(self) -> float:
        with self.lock:
            return self.value

    def sleep(self, seconds: float) -> None:
        with self.lock:
            self.value += max(0.0, seconds)


class UtcClock:
    def __init__(self) -> None:
        self.calls = 0
        self.lock = threading.Lock()

    def __call__(self) -> datetime:
        with self.lock:
            value = TRIGGER + timedelta(milliseconds=self.calls)
            self.calls += 1
            return value


class FakeCamera:
    def __init__(self, *, fail_after_reads: int | None = None, fail_open: bool = False) -> None:
        self.fail_after_reads = fail_after_reads
        self.fail_open = fail_open
        self.read_calls = 0
        self.open_calls = 0
        self.close_calls = 0
        self.shown_states: list[str] = []

    def open(self, settings) -> None:
        del settings
        self.open_calls += 1
        if self.fail_open:
            raise RuntimeError("mock C920 open failure")

    def read(self):
        if self.fail_after_reads is not None and self.read_calls >= self.fail_after_reads:
            return None
        self.read_calls += 1
        return np.full((40, 60, 3), self.read_calls, dtype=np.uint8)

    def show(self, frame, status) -> None:
        del frame
        self.shown_states.append(status)

    def poll_event(self, wait_ms):
        del wait_ms
        from material_level_yolo.camera import PreviewEvent

        return PreviewEvent.NONE

    def close(self) -> None:
        self.close_calls += 1


class FakeDepth:
    def __init__(
        self,
        *,
        fail_capture: bool = False,
        block_capture: threading.Event | None = None,
        capture_started: threading.Event | None = None,
        fail_start: bool = False,
    ) -> None:
        self.fail_capture = fail_capture
        self.block_capture = block_capture
        self.capture_started = capture_started
        self.fail_start = fail_start
        self.start_calls = 0
        self.close_calls = 0
        self.capture_calls = 0
        self.process_calls = 0
        self.reserve_calls = 0
        self.config = SimpleNamespace(
            fusion=SimpleNamespace(burst_frames=45),
            tube=SimpleNamespace(capacity_ml=200.0, usable_height_mm=100.0),
        )

    def start(self) -> None:
        self.start_calls += 1
        if self.fail_start:
            raise RuntimeError("mock D405 start failure")

    def close(self) -> None:
        self.close_calls += 1

    def reserve(self, request):
        self.reserve_calls += 1
        session = Path(request.output_root) / request.experiment_id
        measurements = session / "measurements"
        existing = [path.name for path in measurements.iterdir()] if measurements.is_dir() else []
        index, measurement_id = next_measurement_identity(existing, request.trigger_time_utc)
        measurement = measurements / measurement_id
        depth = measurement / "depth"
        depth.mkdir(parents=True)
        effective_config = session / "effective_config.yaml"
        if not effective_config.exists():
            effective_config.write_text(
                "tube:\n  usable_height_mm: 100.0\n",
                encoding="utf-8",
            )
        manifest = session / "session_manifest.json"
        if manifest.exists():
            data = json.loads(manifest.read_text(encoding="utf-8"))
        else:
            data = {
                "schema_version": 1,
                "experiment_id": request.experiment_id,
                "purpose": request.purpose,
                "measurements": [],
            }
        manifest.write_text(json.dumps(data), encoding="utf-8")
        return SimpleNamespace(
            experiment_id=request.experiment_id,
            measurement_id=measurement_id,
            measurement_index=index,
            trigger_time_utc=format_utc(request.trigger_time_utc),
            session_directory=session,
            measurement_directory=measurement,
            artifact_directory=depth,
            config_sha256=ZERO_HASH,
            calibration_sha256=ONE_HASH,
            software_commit=None,
        )

    def capture(self, request, *, now_utc=None):
        self.capture_calls += 1
        if self.capture_started is not None:
            self.capture_started.set()
        if self.block_capture is not None:
            self.block_capture.wait(timeout=5)
        if self.fail_capture:
            raise RuntimeError("mock D405 burst failure")
        clock = now_utc or (lambda: TRIGGER)
        start = clock()
        end = clock()
        return SimpleNamespace(
            burst=object(),
            trigger_time_utc=format_utc(request.trigger_time_utc),
            capture_start_utc=format_utc(start),
            capture_end_utc=format_utc(end),
        )

    def process(self, captured, request, reservation, *, now_utc=None):
        self.process_calls += 1
        processing = (now_utc or (lambda: TRIGGER))()
        record = prepare_record(
            with_estimate(
                CommonMeasurementRecord(
                    experiment_id=request.experiment_id,
                    measurement_id=reservation.measurement_id,
                    measurement_index=reservation.measurement_index,
                    method="depth",
                    acquisition_mode="synchronized",
                    trigger_time_utc=request.trigger_time_utc,
                    capture_start_utc=captured.capture_start_utc,
                    capture_end_utc=captured.capture_end_utc,
                    processing_time_utc=processing,
                    material_name=request.material_name,
                    total_capacity_ml=200.0,
                    bulk_density_g_per_ml=request.bulk_density_g_per_ml,
                    total_possible_weight_g=request.total_possible_weight_g,
                    reference_material_weight_g=request.reference_material_weight_g,
                    manual_material_level_mm=request.manual_material_level_mm,
                    valid=True,
                    status="complete_valid",
                    artifact_directory=reservation.artifact_directory.relative_to(
                        reservation.session_directory
                    ).as_posix(),
                    source_artifact=(
                        reservation.artifact_directory / "result.json"
                    ).relative_to(reservation.session_directory).as_posix(),
                    config_sha256=ZERO_HASH,
                    calibration_or_model_sha256=ONE_HASH,
                    notes=request.operator_notes,
                ),
                50.0,
            ),
            usable_internal_height_mm=100.0,
        )
        (reservation.artifact_directory / "result.json").write_text("{}", encoding="utf-8")
        MeasurementWorkbookStore(
            reservation.session_directory / "depth_measurements.xlsx", method="depth"
        ).upsert(record, usable_internal_height_mm=100.0)
        session_path = reservation.session_directory / "session_manifest.json"
        session = json.loads(session_path.read_text(encoding="utf-8"))
        session["measurements"] = [
            {
                "measurement_id": reservation.measurement_id,
                "measurement_index": reservation.measurement_index,
                "method": "depth",
                "status": record.status,
                "valid": True,
            }
        ]
        session_path.write_text(json.dumps(session), encoding="utf-8")
        return SimpleNamespace(record=record, estimate=SimpleNamespace())


def experiment_config(tmp_path: Path, *, count: int = 6, span: float = 1.0) -> ExperimentConfig:
    camera = CameraSettings(
        device_index=0,
        backend="auto",
        width=60,
        height=40,
        fps=30.0,
        warmup_frames=0,
        capture_count=6,
        capture_interval_seconds=0.2,
        capture_duration_seconds=None,
        capture_key="c",
        quit_key="q",
        window_name="fake synchronized preview",
        preview_wait_ms=1,
        mirror_preview=False,
        close_on_capture=False,
        freeze_after_capture_ms=0,
    )
    return ExperimentConfig(
        schema_version=1,
        config_path=PROJECT_ROOT / "config.yaml",
        depth_config_path=REPOSITORY_ROOT / "3d_camera" / "config.yaml",
        depth_calibration_path=REPOSITORY_ROOT / "3d_camera" / "calibration" / "empty.npz",
        c920=camera,
        acquisition=AcquisitionSettings(count, span, 2, 0.01),
        output_root=(tmp_path / "Résultats synchronisés avec espaces").resolve(),
        storage_profile="research",
    )


def inputs() -> ExperimentInputs:
    return ExperimentInputs(
        experiment_id="synchronized-study",
        purpose="hardware-free synchronization test",
        material_name="operator polymer",
        total_capacity_ml=200.0,
        bulk_density_g_per_ml=0.5,
        total_possible_weight_g=100.0,
        reference_material_weight_g=25.0,
        manual_material_level_mm=40.0,
        operator_notes="résultat synchronisé",
    )


def build_experiment(tmp_path: Path, camera=None, depth=None, *, count=6, span=1.0):
    clock = VirtualClock()
    camera = camera or FakeCamera()
    depth = depth or FakeDepth()
    states = []
    experiment = SynchronizedExperiment(
        experiment_config(tmp_path, count=count, span=span),
        inputs(),
        c920=camera,
        depth=depth,
        utc_now=UtcClock(),
        monotonic=clock.monotonic,
        sleep=clock.sleep,
        state_callback=states.append,
        software_repository=REPOSITORY_ROOT,
    )
    return experiment, camera, depth, states


def test_fake_trigger_creates_six_distributed_files_and_shared_depth_id(tmp_path: Path) -> None:
    experiment, camera, depth, states = build_experiment(tmp_path)
    summary = experiment.confirmation_summary()
    assert summary["yolo_loaded_or_run"] is False
    experiment.initialize()
    outcome = experiment.trigger()
    assert outcome.final_state is ExperimentState.SAVED
    assert len(outcome.c920.frames) == 6
    assert [item.target_offset_seconds for item in outcome.c920.frames] == pytest.approx(
        [0.0, 0.2, 0.4, 0.6, 0.8, 1.0]
    )
    assert [item.actual_monotonic for item in outcome.c920.frames] == pytest.approx(
        [100.0, 100.2, 100.4, 100.6, 100.8, 101.0]
    )
    assert all(item.path.is_file() and item.path.suffix == ".png" for item in outcome.c920.frames)
    assert outcome.depth_record.measurement_id == outcome.yolo_record.measurement_id == outcome.measurement_id
    assert outcome.yolo_record.status == "pending_offline_inference"
    assert outcome.yolo_record.valid is None
    assert outcome.yolo_record.estimated_material_percent is None
    assert depth.capture_calls == depth.process_calls == 1
    assert ExperimentState.ARMED in states
    assert ExperimentState.CAPTURING in states
    assert ExperimentState.PROCESSING in states
    assert ExperimentState.SAVED in states

    depth_book = load_workbook(outcome.session_directory / "depth_measurements.xlsx")
    yolo_book = load_workbook(outcome.session_directory / "yolo_measurements.xlsx")
    assert (outcome.session_directory / "depth_measurements.csv").is_file()
    assert (outcome.session_directory / "yolo_measurements.csv").is_file()
    depth_columns = tuple(cell.value for cell in depth_book["measurements"][1])
    yolo_columns = tuple(cell.value for cell in yolo_book["measurements"][1])
    assert depth_columns == yolo_columns == COMMON_COLUMNS
    manifest = json.loads(outcome.capture_manifest.read_text(encoding="utf-8"))
    assert manifest["state"] == "SAVED"
    assert manifest["yolo_loaded_or_run"] is False
    assert len(manifest["c920"]["frames"]) == 6
    assert all("timing_error_seconds" in item for item in manifest["c920"]["frames"])
    assert all("target_utc" in item for item in manifest["c920"]["frames"])
    assert manifest["depth"]["capture_start_utc"] is not None
    assert ExperimentState.INITIALIZING in states
    experiment.close()
    assert camera.close_calls == depth.close_calls == 1


def test_automatic_indexes_are_monotonic_across_triggers(tmp_path: Path) -> None:
    experiment, _camera, _depth, _states = build_experiment(tmp_path)
    experiment.initialize()
    first = experiment.trigger()
    second = experiment.trigger()
    assert first.measurement_index == 1
    assert second.measurement_index == 2
    assert first.measurement_id.startswith("000001_taking_")
    assert second.measurement_id.startswith("000002_taking_")
    assert first.measurement_id != second.measurement_id
    experiment.close()


def test_each_trigger_writes_its_own_references_to_matching_depth_and_yolo_rows(
    tmp_path: Path,
) -> None:
    experiment, _camera, _depth, _states = build_experiment(tmp_path)
    experiment.initialize()
    first = experiment.trigger(
        MeasurementReferenceInputs(30.0, 60.0, "first measurement slope")
    )
    second = experiment.trigger(
        MeasurementReferenceInputs(18.0, 25.0, "second measurement flat")
    )

    depth_rows = {
        row.measurement_id: row
        for row in MeasurementWorkbookStore(
            first.session_directory / "depth_measurements.xlsx", method="depth"
        ).load_records()
    }
    yolo_rows = {
        row.measurement_id: row
        for row in MeasurementWorkbookStore(
            first.session_directory / "yolo_measurements.xlsx", method="yolo"
        ).load_records()
    }
    expected = {
        first.measurement_id: (30.0, 60.0, 60.0, 60.0, 120.0, "first measurement slope"),
        second.measurement_id: (18.0, 25.0, 36.0, 25.0, 50.0, "second measurement flat"),
    }
    for measurement_id, values in expected.items():
        weight, level, weight_volume, manual_percent, manual_volume, note = values
        for row in (depth_rows[measurement_id], yolo_rows[measurement_id]):
            assert row.reference_material_weight_g == weight
            assert row.manual_material_level_mm == level
            assert row.reference_volume_from_weight_ml == pytest.approx(weight_volume)
            assert row.manual_material_percent == pytest.approx(manual_percent)
            assert row.reference_volume_from_manual_ml == pytest.approx(manual_volume)
            assert f"Measurement notes: {note}" in row.notes
    first_manifest = json.loads(first.capture_manifest.read_text(encoding="utf-8"))
    assert first_manifest["reference_inputs"] == {
        "reference_material_weight_g": 30.0,
        "manual_material_level_mm": 60.0,
        "measurement_notes": "first measurement slope",
        "general_experiment_notes": "résultat synchronisé",
    }
    assert depth_rows[first.measurement_id].reference_material_weight_g != (
        depth_rows[second.measurement_id].reference_material_weight_g
    )
    experiment.close()


def test_blank_second_trigger_never_reuses_previous_or_legacy_references(
    tmp_path: Path,
) -> None:
    experiment, _camera, _depth, _states = build_experiment(tmp_path)
    experiment.initialize()
    first = experiment.trigger(MeasurementReferenceInputs(12.5, 45.0, "available"))
    second = experiment.trigger(MeasurementReferenceInputs())

    assert first.depth_record.reference_material_weight_g == 12.5
    assert first.yolo_record.manual_material_level_mm == 45.0
    for row in (second.depth_record, second.yolo_record):
        assert row.reference_material_weight_g is None
        assert row.reference_volume_from_weight_ml is None
        assert row.manual_material_level_mm is None
        assert row.manual_material_percent is None
        assert row.reference_volume_from_manual_ml is None
    experiment.close()


def test_cancelled_guided_reference_step_does_not_allocate_or_capture(
    tmp_path: Path,
) -> None:
    experiment, _camera, depth, _states = build_experiment(tmp_path)
    experiment.initialize()
    answers = iter(("54.7", "41.5", "slight slope visible", "c"))
    output: list[str] = []

    outcome = _guided_measurement(
        experiment,
        input_func=lambda _prompt: next(answers),
        print_func=output.append,
    )

    assert outcome is None
    assert experiment.state is ExperimentState.ARMED
    assert experiment.next_measurement_index() == 1
    assert depth.reserve_calls == depth.capture_calls == depth.process_calls == 0
    assert not (experiment.inputs.output_root / experiment.inputs.experiment_id).exists()
    assert any("cancelled before allocation" in line for line in output)
    experiment.close()


def test_guided_blank_reference_inputs_are_saved_as_blanks(tmp_path: Path) -> None:
    experiment, _camera, depth, _states = build_experiment(tmp_path)
    experiment.initialize()
    answers = iter(("", "", "", ""))

    outcome = _guided_measurement(
        experiment,
        input_func=lambda _prompt: next(answers),
        print_func=lambda _line: None,
    )

    assert outcome is not None
    assert depth.reserve_calls == depth.capture_calls == 1
    for row in (outcome.depth_record, outcome.yolo_record):
        assert row.reference_material_weight_g is None
        assert row.reference_volume_from_weight_ml is None
        assert row.manual_material_level_mm is None
        assert row.manual_material_percent is None
        assert row.reference_volume_from_manual_ml is None
    experiment.close()


def test_busy_trigger_is_rejected_without_second_allocation(tmp_path: Path) -> None:
    release = threading.Event()
    started = threading.Event()
    depth = FakeDepth(block_capture=release, capture_started=started)
    experiment, _camera, _depth, _states = build_experiment(tmp_path, depth=depth)
    experiment.initialize()
    results = []
    worker = threading.Thread(target=lambda: results.append(experiment.trigger()))
    worker.start()
    assert started.wait(timeout=2)
    with pytest.raises(BusyTriggerError, match="CAPTURING"):
        experiment.trigger()
    release.set()
    worker.join(timeout=5)
    assert len(results) == 1
    assert len(list((results[0].session_directory / "measurements").iterdir())) == 1
    experiment.close()


def test_partial_c920_is_auditable_and_never_gets_zero_estimates(tmp_path: Path) -> None:
    experiment, _camera, depth, _states = build_experiment(
        tmp_path, camera=FakeCamera(fail_after_reads=3)
    )
    experiment.initialize()
    outcome = experiment.trigger()
    assert outcome.final_state is ExperimentState.PARTIAL
    assert 0 < len(outcome.c920.frames) < 6
    assert outcome.yolo_record.status == "partial_capture"
    assert outcome.yolo_record.valid is None
    assert outcome.yolo_record.estimated_material_volume_ml is None
    assert depth.process_calls == 1
    manifest = json.loads(outcome.capture_manifest.read_text(encoding="utf-8"))
    assert manifest["state"] == "PARTIAL"
    assert manifest["resumable"] is True
    assert manifest["resume_policy"] == "offline_yolo_same_identity"
    assert len(manifest["c920"]["frames"]) == len(outcome.c920.frames)
    for item in manifest["c920"]["frames"]:
        assert (outcome.session_directory / item["path"]).is_file()
    experiment.close()


def test_d405_failure_preserves_c920_and_writes_failed_depth_row(tmp_path: Path) -> None:
    experiment, _camera, depth, _states = build_experiment(
        tmp_path, depth=FakeDepth(fail_capture=True)
    )
    experiment.initialize()
    outcome = experiment.trigger()
    assert outcome.final_state is ExperimentState.PARTIAL
    assert len(outcome.c920.frames) == 6
    assert outcome.depth_record.status == "acquisition_failed"
    assert outcome.depth_record.valid is False
    assert outcome.depth_record.estimated_material_percent is None
    assert outcome.yolo_record.status == "pending_offline_inference"
    assert depth.process_calls == 0
    manifest = json.loads(outcome.capture_manifest.read_text(encoding="utf-8"))
    assert manifest["depth"]["status"] == "failed"
    experiment.close()


def test_initialization_failure_cleans_both_owners_exactly_once(tmp_path: Path) -> None:
    camera = FakeCamera(fail_open=True)
    depth = FakeDepth()
    experiment, _camera, _depth, states = build_experiment(tmp_path, camera=camera, depth=depth)
    with pytest.raises(RuntimeError, match="C920 open"):
        experiment.initialize()
    assert camera.close_calls == depth.close_calls == 1
    experiment.close()
    assert camera.close_calls == depth.close_calls == 1
    assert ExperimentState.ERROR in states
    assert experiment.state is ExperimentState.STOPPED


def test_isolated_orchestrator_import_never_imports_yolo_inference() -> None:
    code = (
        "import sys; "
        f"sys.path[:0]={[str(PROJECT_ROOT), str(REPOSITORY_ROOT), str(REPOSITORY_ROOT / '3d_camera'), str(REPOSITORY_ROOT / 'Material_level_using_yolo')]!r}; "
        "import run_experiment.core; "
        "assert 'material_level_yolo.inference' not in sys.modules; "
        "assert 'material_level_yolo.workflow' not in sys.modules"
    )
    completed = subprocess.run(
        [sys.executable, "-c", code],
        cwd=REPOSITORY_ROOT,
        text=True,
        capture_output=True,
        check=False,
    )
    assert completed.returncode == 0, completed.stderr


def test_fake_end_to_end_capture_then_deferred_yolo_preserves_contract(
    tmp_path: Path,
) -> None:
    """Exercise the public stage boundary without comparing method estimates."""

    class DeferredAdapter:
        def __init__(self) -> None:
            self.mapping = SemanticClassMap(
                material_id=6,
                material_name="polymer",
                empty_id=5,
                empty_name="Empty",
                model_names={5: "Empty", 6: "polymer"},
            )
            self.weights_sha256 = "c" * 64
            self.predict_calls = 0

        def load(self):
            return self.mapping

        def predict(self, image):
            self.predict_calls += 1
            height, width = image.shape[:2]
            boundary = round(height * 0.4)
            material = np.zeros((height, width), dtype=np.float32)
            empty = np.zeros_like(material)
            material[boundary:] = 1.0
            empty[:boundary] = 1.0
            raw = SimpleNamespace(
                boxes=SimpleNamespace(
                    cls=np.asarray([6, 5], dtype=np.float32),
                    conf=np.asarray([0.95, 0.96], dtype=np.float32),
                    xyxy=np.empty((0, 4), dtype=np.float32),
                ),
                masks=SimpleNamespace(data=np.stack([material, empty])),
            )
            return InferenceOutput(
                raw_results=(raw,),
                class_map=self.mapping,
                model_task="segment",
                weights_sha256=self.weights_sha256,
            )

    clock = VirtualClock()
    camera = FakeCamera()
    depth = FakeDepth()
    experiment = SynchronizedExperiment(
        experiment_config(tmp_path),
        replace(
            inputs(),
            experiment_id="publication-ready-fake-study",
            material_name="operator supplied polymer",
            manual_material_level_mm=None,
        ),
        c920=camera,
        depth=depth,
        utc_now=UtcClock(),
        monotonic=clock.monotonic,
        sleep=clock.sleep,
        software_repository=REPOSITORY_ROOT,
    )
    experiment.initialize()
    captured = experiment.trigger(
        MeasurementReferenceInputs(27.5, 35.0, "independent reference")
    )
    experiment.close()
    assert captured.yolo_record.status == "pending_offline_inference"
    assert captured.depth_record.measurement_id == captured.measurement_id
    assert captured.depth_record.reference_material_weight_g == 27.5
    assert captured.yolo_record.manual_material_level_mm == 35.0
    depth_bytes = {
        name: (captured.session_directory / name).read_bytes()
        for name in ("depth_measurements.xlsx", "depth_measurements.csv")
    }

    adapters: list[DeferredAdapter] = []

    def factory(profile):
        assert profile.name == "polymer"
        adapter = DeferredAdapter()
        adapters.append(adapter)
        return adapter

    yolo_project = load_yolo_config(
        REPOSITORY_ROOT / "Material_level_using_yolo" / "config.yaml"
    )
    roi_reference = tmp_path / "calibration" / "tube_roi.json"
    roi_reference.parent.mkdir(parents=True, exist_ok=True)
    roi_reference.write_text("{}\n", encoding="utf-8")
    yolo_project = replace(
        yolo_project,
        profiles={
            name: replace(
                profile,
                tube_roi=replace(
                    profile.tube_roi,
                    frozen=True,
                    left=0.05,
                    right=0.95,
                    calibration_reference=roi_reference,
                ),
            )
            for name, profile in yolo_project.profiles.items()
        },
    )
    processed = process_synchronized_captures(
        yolo_project,
        captured.session_directory,
        adapter_factory=factory,
        now_utc=lambda: TRIGGER + timedelta(seconds=10),
        software_repository=REPOSITORY_ROOT,
    )
    group = processed.groups[0]
    assert group.action == "processed"
    assert group.measurement_id == captured.measurement_id
    assert adapters[0].predict_calls == 6
    yolo_store = MeasurementWorkbookStore(
        captured.session_directory / "yolo_measurements.xlsx", method="yolo"
    )
    rows = yolo_store.load_records()
    assert len(rows) == 1
    yolo_row = rows[0]
    assert yolo_row.measurement_id == captured.measurement_id
    assert yolo_row.status == "complete_valid"
    assert yolo_row.reference_material_weight_g == 27.5
    assert yolo_row.reference_volume_from_weight_ml == pytest.approx(55.0)
    assert yolo_row.manual_material_level_mm == 35.0
    assert yolo_row.manual_material_percent == pytest.approx(35.0)
    assert yolo_row.reference_volume_from_manual_ml == pytest.approx(70.0)
    assert yolo_row.notes == captured.yolo_record.notes
    assert yolo_row.config_sha256 is not None
    assert yolo_row.calibration_or_model_sha256 == "c" * 64
    assert yolo_store.verify_mirror()
    for name, payload in depth_bytes.items():
        assert (captured.session_directory / name).read_bytes() == payload

    depth_book = load_workbook(
        captured.session_directory / "depth_measurements.xlsx", read_only=True
    )
    yolo_book = load_workbook(
        captured.session_directory / "yolo_measurements.xlsx", read_only=True
    )
    try:
        depth_columns = tuple(cell.value for cell in depth_book["measurements"][1])
        yolo_columns = tuple(cell.value for cell in yolo_book["measurements"][1])
        assert depth_columns == yolo_columns == COMMON_COLUMNS
        assert "yolo_image_details" not in depth_book.sheetnames
        assert "yolo_image_details" in yolo_book.sheetnames
    finally:
        depth_book.close()
        yolo_book.close()

    effective = (
        captured.session_directory
        / "provenance"
        / "yolo_effective_config_polymer.yaml"
    )
    assert effective.is_file()
    assert sha256_file(effective) == yolo_row.config_sha256
    artifact_manifest = json.loads(
        (group.artifact_directory / "artifact_manifest.json").read_text(encoding="utf-8")
    )
    assert artifact_manifest["measurement_id"] == captured.measurement_id
    assert all(
        sha256_file(group.artifact_directory / item["path"]) == item["sha256"]
        for item in artifact_manifest["artifacts"]
    )
    capture_manifest = json.loads(captured.capture_manifest.read_text(encoding="utf-8"))
    assert capture_manifest["yolo_loaded_or_run"] is False
    assert capture_manifest["offline_yolo_processing"]["state"] == "COMPLETED"

    def forbidden_factory(_profile):
        raise AssertionError("idempotent second pass must not load YOLO")

    repeated = process_synchronized_captures(
        load_yolo_config(REPOSITORY_ROOT / "Material_level_using_yolo" / "config.yaml"),
        captured.measurement_directory,
        adapter_factory=forbidden_factory,
    )
    assert repeated.groups[0].action == "verified_existing"
    assert len(yolo_store.load_records()) == 1
