from __future__ import annotations

from dataclasses import replace
from datetime import datetime, timedelta, timezone
import os
from pathlib import Path
import re
from unittest.mock import patch

from openpyxl import load_workbook
import pytest

from experiment_records import (
    COMMON_COLUMNS,
    AtomicWriteError,
    CommonMeasurementRecord,
    DuplicateRecordError,
    IdentifierError,
    MeasurementWorkbookStore,
    Method,
    RecordValidationError,
    WorkbookSchemaError,
    capture_provenance,
    format_utc,
    git_revision,
    make_measurement_id,
    next_measurement_identity,
    parse_measurement_id,
    prepare_record,
    sha256_bytes,
    sha256_file,
    validate_experiment_id,
    validate_aligned_records,
    with_estimate,
)


TRIGGER = datetime(2026, 8, 13, 10, 45, 12, 345678, tzinfo=timezone.utc)
ZERO_HASH = "0" * 64
ONE_HASH = "1" * 64
COMMIT = "a" * 40 + "+dirty"


def complete_record(
    *,
    method: str = "depth",
    index: int = 1,
    percent: float = 35.0,
    notes: str | None = None,
) -> CommonMeasurementRecord:
    trigger = TRIGGER + timedelta(seconds=index - 1)
    mode = "synthetic" if method == "depth" else "folder"
    record = CommonMeasurementRecord(
        experiment_id="study-001",
        measurement_id=make_measurement_id(index, trigger),
        measurement_index=index,
        method=method,
        acquisition_mode=mode,
        trigger_time_utc=trigger,
        capture_start_utc=None if mode == "folder" else trigger + timedelta(milliseconds=10),
        capture_end_utc=None if mode == "folder" else trigger + timedelta(milliseconds=510),
        processing_time_utc=trigger + timedelta(seconds=1),
        material_name="Powder" if method == "depth" else "polymer",
        total_capacity_ml=200.0,
        valid=True,
        quality_score=0.9,
        status="complete_valid",
        notes=notes,
        artifact_directory=f"measurements/{make_measurement_id(index, trigger)}/{method}",
        source_artifact=("source image.png" if method == "yolo" else "depth/result.json"),
        config_sha256=ZERO_HASH,
        calibration_or_model_sha256=ONE_HASH,
        software_commit=COMMIT,
    )
    return prepare_record(with_estimate(record, percent))


def pending_yolo_record(index: int = 1) -> CommonMeasurementRecord:
    trigger = TRIGGER + timedelta(seconds=index - 1)
    return prepare_record(
        CommonMeasurementRecord(
            experiment_id="study-001",
            measurement_id=make_measurement_id(index, trigger),
            measurement_index=index,
            method="yolo",
            acquisition_mode="synchronized",
            trigger_time_utc=trigger,
            capture_start_utc=trigger + timedelta(milliseconds=10),
            capture_end_utc=trigger + timedelta(seconds=1, milliseconds=10),
            total_capacity_ml=200.0,
            valid=None,
            status="pending_offline_inference",
            artifact_directory=f"measurements/{make_measurement_id(index, trigger)}/yolo",
            source_artifact="capture_manifest.json",
        )
    )


class TestIdentifiers:
    def test_measurement_id_is_sortable_and_uses_truncated_utc_milliseconds(self) -> None:
        value = make_measurement_id(1, TRIGGER)
        assert value == "000001_taking_20260813T104512345Z"
        parsed = parse_measurement_id(value)
        assert parsed.index == 1
        assert parsed.trigger_time_utc.microsecond == 345000

    def test_offset_timestamp_is_normalized_before_id_creation(self) -> None:
        offset = TRIGGER.astimezone(timezone(timedelta(hours=2)))
        assert make_measurement_id(1, offset) == make_measurement_id(1, TRIGGER)
        assert format_utc(offset) == "2026-08-13T10:45:12.345678Z"

    def test_naive_time_and_unstable_ids_are_rejected(self) -> None:
        with pytest.raises(IdentifierError, match="timezone"):
            make_measurement_id(1, datetime(2026, 8, 13, 10, 0, 0))
        with pytest.raises(IdentifierError, match="timestamp"):
            prepare_record(
                replace(
                    complete_record(),
                    measurement_id="000001_taking_20260813T104512344Z",
                )
            )

    @pytest.mark.parametrize(
        "value", ["", "has space", "../escape", "CON", "NUL.txt", "éxperiment"]
    )
    def test_nonportable_experiment_ids_are_rejected(self, value: str) -> None:
        with pytest.raises(IdentifierError):
            validate_experiment_id(value)

    def test_next_identity_never_reuses_an_existing_index(self) -> None:
        first = make_measurement_id(1, TRIGGER)
        third = make_measurement_id(3, TRIGGER + timedelta(seconds=2))
        index, value = next_measurement_identity(
            [first, third], TRIGGER + timedelta(seconds=3)
        )
        assert index == 4
        assert value.startswith("000004_taking_")


class TestValidationAndFormulas:
    def test_explicit_reference_inputs_are_derived_and_density_is_retained(self) -> None:
        base = replace(
            complete_record(),
            bulk_density_g_per_ml=0.5,
            total_possible_weight_g=100.0,
            reference_material_weight_g=25.0,
            reference_volume_from_weight_ml=None,
            manual_material_level_mm=40.0,
            manual_material_percent=None,
            reference_volume_from_manual_ml=None,
        )
        prepared = prepare_record(base, usable_internal_height_mm=80.0)
        assert prepared.bulk_density_g_per_ml == 0.5
        assert prepared.reference_volume_from_weight_ml == 50.0
        assert prepared.manual_material_percent == 50.0
        assert prepared.reference_volume_from_manual_ml == 100.0

    def test_material_name_never_infers_density_or_weight_reference(self) -> None:
        prepared = prepare_record(replace(complete_record(), material_name="known powder"))
        assert prepared.bulk_density_g_per_ml is None
        assert prepared.reference_volume_from_weight_ml is None

    @pytest.mark.parametrize(
        ("field", "value", "message"),
        [
            ("total_capacity_ml", float("nan"), "finite"),
            ("reference_material_weight_g", -0.1, "nonnegative"),
            ("quality_score", 1.1, r"\[0, 1\]"),
            ("estimated_material_percent", 101.0, r"\[0, 100\]"),
        ],
    )
    def test_invalid_numeric_values_are_rejected(
        self, field: str, value: float, message: str
    ) -> None:
        with pytest.raises(RecordValidationError, match=message):
            prepare_record(replace(complete_record(), **{field: value}))

    def test_timestamps_must_be_chronological(self) -> None:
        with pytest.raises(RecordValidationError, match="chronological"):
            prepare_record(
                replace(
                    complete_record(),
                    capture_start_utc=TRIGGER + timedelta(seconds=2),
                    capture_end_utc=TRIGGER + timedelta(seconds=1),
                )
            )

    def test_valid_result_must_conserve_percent_and_capacity(self) -> None:
        with pytest.raises(RecordValidationError, match="percentages"):
            prepare_record(replace(complete_record(), estimated_empty_percent=64.0))
        with pytest.raises(RecordValidationError, match="volumes"):
            prepare_record(replace(complete_record(), estimated_empty_volume_ml=120.0))

    def test_invalid_result_keeps_estimates_blank_not_zero(self) -> None:
        invalid = prepare_record(
            replace(
                complete_record(),
                valid=False,
                status="complete_invalid",
                estimated_material_percent=None,
                estimated_material_volume_ml=None,
                estimated_empty_percent=None,
                estimated_empty_volume_ml=None,
            )
        )
        for name in (
            "estimated_material_percent",
            "estimated_material_volume_ml",
            "estimated_empty_percent",
            "estimated_empty_volume_ml",
        ):
            assert invalid.to_mapping()[name] is None

    def test_reference_volume_cannot_exist_without_both_explicit_inputs(self) -> None:
        with pytest.raises(RecordValidationError, match="explicit weight and density"):
            prepare_record(replace(complete_record(), reference_volume_from_weight_ml=1.0))

    def test_zero_density_and_zero_manual_height_are_validation_errors(self) -> None:
        with pytest.raises(RecordValidationError, match="greater than zero"):
            prepare_record(
                replace(
                    complete_record(),
                    bulk_density_g_per_ml=0.0,
                    reference_material_weight_g=1.0,
                )
            )
        with pytest.raises(RecordValidationError, match="greater than zero"):
            prepare_record(
                replace(complete_record(), manual_material_level_mm=1.0),
                usable_internal_height_mm=0.0,
            )

    def test_pending_record_cannot_claim_a_processing_time(self) -> None:
        with pytest.raises(RecordValidationError, match="must be blank"):
            prepare_record(
                replace(
                    pending_yolo_record(),
                    processing_time_utc=TRIGGER + timedelta(seconds=2),
                )
            )

    def test_aligned_records_require_identical_shared_metadata_only(self) -> None:
        depth = complete_record(method="depth")
        yolo = replace(
            complete_record(method="yolo", percent=80.0),
            material_name=depth.material_name,
            notes=depth.notes,
        )
        aligned_depth, aligned_yolo = validate_aligned_records(depth, yolo)
        assert aligned_depth.estimated_material_percent == 35.0
        assert aligned_yolo.estimated_material_percent == 80.0
        with pytest.raises(RecordValidationError, match="total_capacity_ml"):
            validate_aligned_records(
                depth,
                replace(
                    yolo,
                    total_capacity_ml=250.0,
                    estimated_material_volume_ml=200.0,
                    estimated_empty_volume_ml=50.0,
                ),
            )


class TestWorkbookStore:
    def test_code_schema_matches_the_frozen_data_dictionary(self) -> None:
        dictionary = (
            Path(__file__).resolve().parents[1]
            / "docs"
            / "material_level_study"
            / "DATA_DICTIONARY.md"
        ).read_text(encoding="utf-8")
        documented = tuple(
            match.group(1)
            for match in re.finditer(r"(?m)^\| [0-9]+ \| `([^`]+)` \|", dictionary)
        )
        assert documented == COMMON_COLUMNS

    def test_unicode_paths_blanks_and_column_parity(self, tmp_path: Path) -> None:
        root = tmp_path / "étude avec espaces"
        depth_path = root / "depth_measurements.xlsx"
        yolo_path = root / "yolo_measurements.xlsx"
        depth_store = MeasurementWorkbookStore(depth_path, method=Method.DEPTH)
        yolo_store = MeasurementWorkbookStore(yolo_path, method=Method.YOLO)
        depth_store.create([complete_record()])
        yolo_store.create([pending_yolo_record()])

        depth_sheet = load_workbook(depth_path)["measurements"]
        yolo_sheet = load_workbook(yolo_path)["measurements"]
        depth_headers = tuple(cell.value for cell in depth_sheet[1])
        yolo_headers = tuple(cell.value for cell in yolo_sheet[1])
        assert depth_headers == yolo_headers == COMMON_COLUMNS
        assert len(depth_headers) == 32
        density_column = COMMON_COLUMNS.index("bulk_density_g_per_ml") + 1
        assert depth_sheet.cell(row=2, column=density_column).value is None
        assert depth_store.verify_mirror()
        assert yolo_store.verify_mirror()
        csv_text = depth_path.with_suffix(".csv").read_text(encoding="utf-8-sig")
        assert "None" not in csv_text and "NaN" not in csv_text

    def test_create_append_and_idempotent_upsert(self, tmp_path: Path) -> None:
        store = MeasurementWorkbookStore(tmp_path / "depth.xlsx", method="depth")
        first = complete_record(index=1)
        second = complete_record(index=2)
        store.create([first])
        store.append(second)
        assert len(store.load_records()) == 2
        with pytest.raises(DuplicateRecordError):
            store.append(first)

        revised = replace(second, notes="operator revision")
        store.upsert(revised)
        store.upsert(revised)
        loaded = store.load_records()
        assert len(loaded) == 2
        assert loaded[1].notes == "operator revision"
        assert store.verify_mirror()

    def test_common_row_and_method_diagnostics_upsert_together(self, tmp_path: Path) -> None:
        store = MeasurementWorkbookStore(tmp_path / "yolo.xlsx", method="yolo")
        record = complete_record(method="yolo")
        columns = ("experiment_id", "measurement_id", "image_index", "valid")
        first_rows = (
            (record.experiment_id, record.measurement_id, 1, True),
            (record.experiment_id, record.measurement_id, 2, False),
        )
        key = {
            "experiment_id": record.experiment_id,
            "measurement_id": record.measurement_id,
        }
        store.upsert_with_diagnostic_rows(
            record,
            sheet_name="yolo_image_details",
            columns=columns,
            rows=first_rows,
            replace_where=key,
        )
        revised = replace(record, notes="rerun diagnostics")
        replacement = ((record.experiment_id, record.measurement_id, 1, True),)
        store.upsert_with_diagnostic_rows(
            revised,
            sheet_name="yolo_image_details",
            columns=columns,
            rows=replacement,
            replace_where=key,
        )
        workbook = load_workbook(store.workbook_path, data_only=True)
        sheet = workbook["yolo_image_details"]
        assert tuple(cell.value for cell in sheet[1]) == columns
        assert tuple(cell.value for cell in sheet[2]) == replacement[0]
        assert sheet.max_row == 2
        assert store.load_records()[0].notes == "rerun diagnostics"
        assert store.verify_mirror()

    def test_failed_second_replace_is_recovered_by_next_upsert(self, tmp_path: Path) -> None:
        store = MeasurementWorkbookStore(tmp_path / "depth.xlsx", method="depth")
        store.create([complete_record(index=1)])
        second = complete_record(index=2)
        real_replace = os.replace
        failed = False

        def fail_csv_once(source: str | Path, destination: str | Path) -> None:
            nonlocal failed
            if Path(destination) == store.csv_path and not failed:
                failed = True
                raise OSError("simulated interruption")
            real_replace(source, destination)

        with patch("experiment_records.export.os.replace", side_effect=fail_csv_once):
            with pytest.raises(AtomicWriteError, match="reconcile"):
                store.upsert(second)

        assert len(store.load_records()) == 2
        assert not store.verify_mirror()
        store.upsert(second)
        assert store.verify_mirror()
        assert len(store.load_records()) == 2
        assert not list(tmp_path.glob("*.tmp.*"))
        assert not store.lock_path.exists()

    def test_store_rejects_a_record_from_the_other_method(self, tmp_path: Path) -> None:
        store = MeasurementWorkbookStore(tmp_path / "depth.xlsx", method="depth")
        with pytest.raises(WorkbookSchemaError, match="cannot contain"):
            store.create([pending_yolo_record()])


class TestProvenance:
    def test_hashes_exact_bytes_and_git_unavailable_is_nonfatal(self, tmp_path: Path) -> None:
        config = tmp_path / "effective config.yaml"
        model = tmp_path / "model weights.bin"
        config.write_bytes(b"alpha: 1\n")
        model.write_bytes(b"weights")
        assert sha256_file(config) == sha256_bytes(b"alpha: 1\n")
        with patch(
            "experiment_records.provenance.subprocess.run", side_effect=OSError("no git")
        ):
            provenance = capture_provenance(
                config_path=config,
                calibration_or_model_path=model,
                repository=tmp_path,
            )
        assert provenance.software_commit is None
        assert provenance.config_sha256 == sha256_bytes(b"alpha: 1\n")
        assert provenance.calibration_or_model_sha256 == sha256_bytes(b"weights")

    def test_current_repository_revision_reports_commit_with_optional_dirty_marker(self) -> None:
        repository = Path(__file__).resolve().parents[1]
        revision = git_revision(repository)
        assert revision is not None
        commit = revision.removesuffix("+dirty")
        assert len(commit) == 40
        assert all(character in "0123456789abcdef" for character in commit)
