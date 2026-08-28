"""Atomic CSV/XLSX/JSON outputs for numerical results and traceability."""

from __future__ import annotations

import csv
import json
import os
from pathlib import Path
import tempfile
from typing import Any, Iterable

from openpyxl import Workbook

from .analysis import AnalysisResult
from .config import AnalysisConfig


SUMMARY_COLUMNS = ("method", "mae_ml", "rmse_ml", "mean_bias_ml", "error_sd_ml", "r_squared", "valid_measurement_count")
TRACE_COLUMNS = ("experiment_id", "measurement_id", "measurement_index", "method", "reference_volume_ml", "estimate_volume_ml", "error_ml", "absolute_error_ml", "reference_source", "source_status")
EXCLUDED_COLUMNS = ("measurement_id", "method", "reason")


def export_tables(result: AnalysisResult, config: AnalysisConfig) -> list[Path]:
    output = config.output_directory
    output.mkdir(parents=True, exist_ok=True)
    paths = [
        _csv(output / "summary_statistics.csv", SUMMARY_COLUMNS, result.summary),
        _csv(output / "matched_measurements.csv", TRACE_COLUMNS, (_as_dict(item) for item in result.comparisons)),
        _csv(output / "excluded_comparisons.csv", EXCLUDED_COLUMNS, result.excluded),
        _xlsx(output / "analysis_results.xlsx", result),
    ]
    report = {
        "experiment_directory": str(config.experiment_directory),
        "input_files": {method: str(path) for method, path in result.input_files.items()},
        "valid_comparison_counts": {row["method"]: row["valid_measurement_count"] for row in result.summary},
        "reference_source_counts": result.reference_source_counts,
        "full_mass_fraction_fallback_enabled": config.allow_full_mass_fraction_fallback,
        "full_mass_fraction_formula": "total_capacity_ml * reference_material_weight_g / total_possible_weight_g",
        "excluded_comparison_count": len(result.excluded),
        "missing_values_invented": False,
    }
    paths.append(_json(output / "analysis_report.json", report))
    return paths


def _as_dict(item) -> dict[str, Any]:
    return {name: getattr(item, name) for name in TRACE_COLUMNS}


def _csv(path: Path, columns: tuple[str, ...], rows: Iterable[dict[str, Any]]) -> Path:
    temp = _temporary(path)
    try:
        with temp.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=columns, extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                writer.writerow({key: "" if row.get(key) is None else row.get(key) for key in columns})
        os.replace(temp, path)
    finally:
        temp.unlink(missing_ok=True)
    return path


def _xlsx(path: Path, result: AnalysisResult) -> Path:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "summary"
    _sheet(summary, SUMMARY_COLUMNS, result.summary)
    matched = workbook.create_sheet("matched_measurements")
    _sheet(matched, TRACE_COLUMNS, (_as_dict(item) for item in result.comparisons))
    excluded = workbook.create_sheet("excluded_comparisons")
    _sheet(excluded, EXCLUDED_COLUMNS, result.excluded)
    temp = _temporary(path)
    try:
        workbook.save(temp)
        os.replace(temp, path)
    finally:
        workbook.close()
        temp.unlink(missing_ok=True)
    return path


def _sheet(sheet, columns: tuple[str, ...], rows: Iterable[dict[str, Any]]) -> None:
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(column) for column in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _json(path: Path, data: dict[str, Any]) -> Path:
    temp = _temporary(path)
    try:
        temp.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        os.replace(temp, path)
    finally:
        temp.unlink(missing_ok=True)
    return path


def _temporary(path: Path) -> Path:
    descriptor, name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    os.close(descriptor)
    return Path(name)

