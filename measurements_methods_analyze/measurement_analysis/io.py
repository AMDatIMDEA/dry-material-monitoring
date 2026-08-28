"""Read common measurement tables without importing or mutating source projects."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REQUIRED_COLUMNS = {
    "experiment_id", "measurement_id", "measurement_index", "method", "status",
    "valid", "total_capacity_ml", "bulk_density_g_per_ml",
    "total_possible_weight_g", "reference_material_weight_g",
    "reference_volume_from_weight_ml", "manual_material_percent",
    "reference_volume_from_manual_ml", "estimated_material_volume_ml",
}


def discover_measurement_files(experiment: Path, preference: str) -> dict[str, Path]:
    if not experiment.is_dir():
        raise FileNotFoundError(f"Experiment directory does not exist: {experiment}")
    order = (preference, "xlsx" if preference == "csv" else "csv")
    found: dict[str, Path] = {}
    for method in ("depth", "yolo"):
        for suffix in order:
            candidate = experiment / f"{method}_measurements.{suffix}"
            if candidate.is_file():
                found[method] = candidate
                break
    if not found:
        raise FileNotFoundError(
            f"No depth_measurements or yolo_measurements CSV/XLSX files found in {experiment}"
        )
    return found


def read_measurement_file(path: Path) -> list[dict[str, Any]]:
    if path.suffix.casefold() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
    elif path.suffix.casefold() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if "measurements" not in workbook.sheetnames:
                raise ValueError(f"Workbook has no measurements sheet: {path}")
            values = workbook["measurements"].iter_rows(values_only=True)
            headers = tuple(next(values, ()))
            rows = [dict(zip(headers, row, strict=True)) for row in values]
        finally:
            workbook.close()
    else:
        raise ValueError(f"Unsupported input format: {path}")
    if not rows:
        return []
    missing = REQUIRED_COLUMNS - set(rows[0])
    if missing:
        raise ValueError(f"{path.name} is missing columns: {', '.join(sorted(missing))}")
    return rows

