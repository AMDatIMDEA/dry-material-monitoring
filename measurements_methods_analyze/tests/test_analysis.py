from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import load_workbook, Workbook
import pytest
import yaml

from measurement_analysis.analysis import analyze_experiment
from measurement_analysis.cli import main
from measurement_analysis.config import load_config
from measurement_analysis.io import discover_measurement_files
from measurement_analysis.plots import _measurement_sequence_plot


HEADERS = (
    "experiment_id", "measurement_id", "measurement_index", "method", "status",
    "valid", "total_capacity_ml", "bulk_density_g_per_ml",
    "total_possible_weight_g", "reference_material_weight_g",
    "reference_volume_from_weight_ml", "manual_material_percent",
    "reference_volume_from_manual_ml", "estimated_material_volume_ml",
)


def row(index: int, method: str, *, status: str = "complete_valid", valid="true", reference_mass=50, full_mass=100, reference_volume="", manual_volume=100, estimate=100):
    return {
        "experiment_id": "study-1",
        "measurement_id": f"{index:06d}_taking_20260101T00000{index}000Z",
        "measurement_index": index,
        "method": method,
        "status": status,
        "valid": valid,
        "total_capacity_ml": 200,
        "bulk_density_g_per_ml": "",
        "total_possible_weight_g": full_mass,
        "reference_material_weight_g": reference_mass,
        "reference_volume_from_weight_ml": reference_volume,
        "manual_material_percent": 50,
        "reference_volume_from_manual_ml": manual_volume,
        "estimated_material_volume_ml": estimate,
    }


def write_csv(path: Path, rows) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=HEADERS)
        writer.writeheader()
        writer.writerows(rows)


def write_config(path: Path, experiment: Path, output: str = "outputs", **reference) -> Path:
    path.write_text(
        yaml.safe_dump(
            {
                "experiment_directory": str(experiment),
                "output_directory": output,
                "input_preference": "csv",
                "reference": {"allow_full_mass_fraction_fallback": reference.get("fallback", True)},
                "repeatability": {"grouping_tolerance_ml": 0.25},
                "plots": {"dpi": 100, "formats": ["png", "pdf"]},
            },
            sort_keys=False,
        ),
        encoding="utf-8",
    )
    return path


def test_matches_ids_uses_gravimetric_fallback_and_skips_invalid(tmp_path: Path) -> None:
    experiment = tmp_path / "expériment data"
    depth = [
        row(1, "depth", estimate=110),
        row(2, "depth", reference_mass=25, manual_volume=55, estimate=48),
    ]
    yolo = [
        row(1, "yolo", manual_volume=100.0000000001, estimate=90),
        row(2, "yolo", reference_mass=25, manual_volume=55, status="complete_invalid", valid="false", estimate=""),
    ]
    write_csv(experiment / "depth_measurements.csv", depth)
    write_csv(experiment / "yolo_measurements.csv", yolo)
    config = load_config(write_config(tmp_path / "config.yaml", experiment))
    result = analyze_experiment(config)

    assert len([item for item in result.comparisons if item.method == "Human"]) == 2
    assert len([item for item in result.comparisons if item.method == "D405"]) == 2
    assert len([item for item in result.comparisons if item.method == "YOLO"]) == 1
    first_depth = next(item for item in result.comparisons if item.method == "D405" and item.measurement_index == 1)
    assert first_depth.reference_volume_ml == 100
    assert first_depth.error_ml == 10
    assert first_depth.measurement_id.startswith("000001_")
    assert any(item["method"] == "YOLO" and "complete_invalid" in item["reason"] for item in result.excluded)
    summary = {item["method"]: item for item in result.summary}
    assert summary["D405"]["mae_ml"] == pytest.approx(6.0)
    assert summary["D405"]["valid_measurement_count"] == 2

    figure = _measurement_sequence_plot(result, config)
    try:
        yolo_line = next(line for line in figure.axes[0].lines if line.get_label() == "YOLO Estimation")
        yolo_values = list(yolo_line.get_ydata())
        assert yolo_values[0] == 90
        assert str(yolo_values[1]) == "nan"
    finally:
        figure.clear()


def test_explicit_reference_volume_has_priority_and_missing_is_not_invented(tmp_path: Path) -> None:
    experiment = tmp_path / "experiment"
    write_csv(experiment / "depth_measurements.csv", [row(1, "depth", reference_volume=80, reference_mass="", full_mass="", estimate=82)])
    config = load_config(write_config(tmp_path / "config.yaml", experiment, fallback=False))
    result = analyze_experiment(config)
    assert {item.reference_volume_ml for item in result.comparisons} == {80.0}
    assert result.reference_source_counts == {"reference_volume_from_weight_ml": 1}

    write_csv(experiment / "depth_measurements.csv", [row(1, "depth", reference_volume="", reference_mass="", full_mass="", estimate=82)])
    missing = analyze_experiment(config)
    assert not missing.comparisons
    assert all(item["reason"] == "gravimetric_reference_unavailable" for item in missing.excluded)


def test_xlsx_fallback_is_discovered(tmp_path: Path) -> None:
    experiment = tmp_path / "experiment"
    experiment.mkdir()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "measurements"
    sheet.append(HEADERS)
    sheet.append(tuple(row(1, "depth")[key] for key in HEADERS))
    workbook.save(experiment / "depth_measurements.xlsx")
    workbook.close()
    assert discover_measurement_files(experiment, "csv")["depth"].suffix == ".xlsx"


def test_cli_exports_all_figures_tables_and_traceability(tmp_path: Path) -> None:
    experiment = tmp_path / "experiment"
    depth_rows, yolo_rows = [], []
    for index, mass in enumerate((100, 75, 75, 50, 25, 0), 1):
        reference = 2 * mass
        depth_rows.append(row(index, "depth", reference_mass=mass, manual_volume=reference + 1, estimate=reference + 2))
        yolo_rows.append(row(index, "yolo", reference_mass=mass, manual_volume=reference + 1, estimate=reference - 3))
    write_csv(experiment / "depth_measurements.csv", depth_rows)
    write_csv(experiment / "yolo_measurements.csv", yolo_rows)
    config_path = write_config(tmp_path / "analysis config.yaml", experiment, "résultats outputs")

    assert main(["--config", str(config_path)]) == 0
    output = tmp_path / "résultats outputs"
    expected = {
        "summary_statistics.csv", "matched_measurements.csv", "excluded_comparisons.csv",
        "analysis_results.xlsx", "analysis_report.json",
    }
    expected.update({f"0{index}_{name}.{suffix}" for index, name in (
        (1, "estimated_vs_reference_volume"),
        (2, "signed_error_vs_reference_fill"),
        (3, "absolute_error_distribution"),
        (4, "bland_altman_agreement"),
        (5, "repeatability_by_reference_fill"),
    ) for suffix in ("png", "pdf")})
    assert expected <= {path.name for path in output.iterdir()}
    with (output / "matched_measurements.csv").open(encoding="utf-8-sig") as stream:
        matched = list(csv.DictReader(stream))
    assert len(matched) == 18
    assert all(item["measurement_id"] for item in matched)
    report = json.loads((output / "analysis_report.json").read_text(encoding="utf-8"))
    assert report["missing_values_invented"] is False
    workbook = load_workbook(output / "analysis_results.xlsx", read_only=True)
    assert workbook.sheetnames == ["summary", "matched_measurements", "excluded_comparisons"]
    workbook.close()
    assert (output / "06_measurement_by_measurement_volume.png").is_file()
    assert (output / "06_measurement_by_measurement_volume.pdf").is_file()
